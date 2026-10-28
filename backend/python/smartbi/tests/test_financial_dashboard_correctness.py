"""财务看板口径回归测试 (2026-10-21, 三轮 Codex 审查复现)。

每个用例对应一条已复现的错算 / 错显, 断言读的是用户看得见的 KPI 与明细表,
而不是实现细节: 修复必须让页面上的数字变对, 不只是让某个内部字段变对。

Usage:
    cd backend/python
    python -m pytest smartbi/tests/test_financial_dashboard_correctness.py -v
"""
import io
import os
import re
import sys
from typing import Any, Dict, List, Optional

import pandas as pd
import pytest

_project_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

from smartbi.services.financial_dashboard import FinancialDashboardService  # noqa: E402
from smartbi.services.financial.registry import registry  # noqa: E402

svc = FinancialDashboardService()


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def build(chart_type: str, rows: List[Dict[str, Any]], start: int = 1, end: int = 1,
          period_type: str = 'month_range', year: int = 2026) -> Dict[str, Any]:
    return svc.generate_chart(chart_type, pd.DataFrame(rows), year, period_type, start, end)


_SCALE = {'亿': 1e8, '万': 1e4}


def to_yuan(text: Any) -> Optional[float]:
    """把 KPI 显示值 ("1,100.00" / "10.00万" / "-") 还原成元; 无值返回 None。"""
    if text is None:
        return None
    s = str(text).strip()
    if s in ('', '-', '—'):
        return None
    m = re.fullmatch(r'([+-]?[\d,]+(?:\.\d+)?)\s*(亿|万)?', s)
    assert m, f'KPI 值不是数字: {text!r}'
    return float(m.group(1).replace(',', '')) * _SCALE.get(m.group(2) or '', 1)


def kpi(result: Dict[str, Any], label: str) -> Dict[str, Any]:
    for k in result.get('kpis') or []:
        if k.get('label') == label:
            return k
    raise AssertionError(f'没有 KPI {label!r}; 实际: {[k.get("label") for k in result.get("kpis") or []]}')


def table_row(result: Dict[str, Any], label: str) -> List[Any]:
    for row in (result.get('tableData') or {}).get('rows') or []:
        if isinstance(row, dict) and row.get('label') == label:
            return row['values']
    raise AssertionError(f'明细表没有 {label!r} 行')


# ---------------------------------------------------------------------------
# 第一轮
# ---------------------------------------------------------------------------
class TestChannelNeverFabricated:
    def test_no_channel_column_reports_missing_instead_of_inventing_channels(self):
        r = build('channel_analysis', [{'月份': '1月', '品类': '酱料', '实际': 100}])
        labels = {str(k.get('value')) for k in r.get('kpis') or []}
        assert not labels & {'直营', '经销', '电商'}, f'缺渠道列时编造了渠道: {r.get("kpis")}'
        assert r['success'] is False
        assert r['metadata']['dataQuality'] == 'insufficient'
        assert '渠道' in r['analysisContext']

    def test_real_channel_column_is_summed_exactly(self):
        r = build('channel_analysis', [
            {'月份': '1月', '渠道': '商超', '实际': 300},
            {'月份': '1月', '渠道': '电商', '实际': 100},
        ])
        assert r['success'] is True
        assert to_yuan(kpi(r, '总收入')['value']) == 400
        assert kpi(r, '最大渠道')['value'] == '商超'


class TestGrossMarginWeighting:
    def test_cumulative_margin_is_total_profit_over_total_revenue(self):
        # 1 月 收入 100 成本 50 (50%), 2 月 收入 900 成本 810 (10%)
        # 正确累计 = (1000-860)/1000 = 14%; 简单平均会给 30%
        r = build('gross_margin_trend', [
            {'月份': '1月', '项目': '营业收入', '实际': 100, '上年': 100},
            {'月份': '1月', '项目': '营业成本', '实际': 50, '上年': 60},
            {'月份': '2月', '项目': '营业收入', '实际': 900, '上年': 100},
            {'月份': '2月', '项目': '营业成本', '实际': 810, '上年': 60},
        ], 1, 2)
        assert float(kpi(r, '本年累计毛利率')['value']) == pytest.approx(14.0, abs=0.05)


class TestYearIsolation:
    def test_single_year_excludes_other_years(self):
        rows = [
            {'月份': '2025-01', '预算': 500, '实际': 400},
            {'月份': '2026-01', '预算': 200, '实际': 100},
        ]
        r = build('budget_achievement', rows, 1, 12, period_type='year', year=2026)
        assert to_yuan(kpi(r, '年度实际')['value']) == 100
        assert to_yuan(kpi(r, '年度目标')['value']) == 200

    def test_single_month_only_that_month(self):
        rows = [
            {'月份': '1月', '预算': 100, '实际': 90},
            {'月份': '3月', '预算': 300, '实际': 330},
        ]
        r = build('budget_achievement', rows, 3, 3, period_type='month')
        assert to_yuan(kpi(r, '年度实际')['value']) == 330


# ---------------------------------------------------------------------------
# 第二轮
# ---------------------------------------------------------------------------
PNL_REFERENCE = [
    ('营业收入', 8333), ('营业成本', 1471), ('销售费用', 540), ('管理费用', 329),
    ('研发费用', 999), ('财务费用', 604), ('其他收益', 65), ('营业外收入', 42),
    ('营业外支出', 86), ('所得税', 1103), ('净利润', 3308),
]


class TestPnlWaterfall:
    def test_net_profit_includes_other_gains_and_non_operating_items(self):
        r = build('pnl_waterfall', [{'月份': '1月', '项目': k, '实际': v} for k, v in PNL_REFERENCE])
        assert to_yuan(kpi(r, '净利润')['value']) == pytest.approx(3308)
        assert not r['metadata'].get('reconciliation'), r['metadata'].get('reconciliation')

    def test_reported_net_profit_mismatch_is_surfaced(self):
        items = [(k, 3300 if k == '净利润' else v) for k, v in PNL_REFERENCE]
        r = build('pnl_waterfall', [{'月份': '1月', '项目': k, '实际': v} for k, v in items])
        notes = r['metadata'].get('reconciliation') or []
        assert any('净利润' in n and '8' in n for n in notes), notes

    def test_subtotal_and_its_details_are_not_double_counted(self):
        r = build('pnl_waterfall', [{'月份': '1月', '项目': k, '实际': v} for k, v in [
            ('营业收入', 1000), ('主营业务收入', 800), ('其他业务收入', 200), ('营业成本', 400)]])
        assert to_yuan(kpi(r, '营业收入')['value']) == 1000

    def test_negative_finance_expense_reduces_expenses(self):
        # 财务费用为负 = 利息收入大于支出, 应增加利润
        r = build('pnl_waterfall', [{'月份': '1月', '项目': k, '实际': v} for k, v in [
            ('营业收入', 1000), ('营业成本', 400), ('管理费用', 100), ('财务费用', -20)]])
        assert to_yuan(kpi(r, '净利润')['value']) == pytest.approx(520)


class TestPerformanceChartsUseRevenueOnly:
    """业绩类图表 (预算完成 / 方差 / 目标达成): 损益表里只取收入, 不把成本费用一起加进「实际」。"""

    PNL_MONTHLY = [
        {'月份': '1月', '项目': '营业收入', '预算': 1000, '实际': 900, '上年': 800},
        {'月份': '1月', '项目': '营业成本', '预算': 600, '实际': 550, '上年': 500},
        {'月份': '1月', '项目': '销售费用', '预算': 100, '实际': 120, '上年': 90},
        {'月份': '3月', '项目': '营业收入', '预算': 1200, '实际': 1300, '上年': 1000},
        {'月份': '3月', '项目': '营业成本', '预算': 700, '实际': 760, '上年': 600},
    ]

    def test_budget_achievement_counts_revenue_only(self):
        r = build('budget_achievement', self.PNL_MONTHLY, 1, 3)
        assert to_yuan(kpi(r, '年度实际')['value']) == 2200
        assert to_yuan(kpi(r, '年度目标')['value']) == 2200
        assert float(kpi(r, '年度达成率')['value']) == pytest.approx(100.0)

    def test_budget_achievement_missing_month_is_empty(self):
        r = build('budget_achievement', self.PNL_MONTHLY, 1, 3)
        assert table_row(r, '实际')[1] is None
        assert table_row(r, '达成率(%)')[1] is None

    def test_variance_analysis_counts_revenue_only(self):
        r = build('variance_analysis', self.PNL_MONTHLY, 1, 3)
        assert to_yuan(kpi(r, '总实际')['value']) == 2200


class TestSheetSemantics:
    """客户样图验收中暴露: 比率表进了金额图、费用表被当成业绩、记分卡把收入显示成 0。"""

    MARGIN_SHEET = [{'月份': f'{m}月', '本年毛利率': f'{25 + m}%', '上年毛利率': f'{24 + m}%'} for m in range(1, 4)]
    EXPENSE_SHEET = [{'月份': f'{m}月', '预算（万元）': 100, '实际（万元）': 90, '上年（万元）': 80} for m in range(1, 4)]

    def test_ratio_sheet_only_feeds_ratio_charts(self):
        r = svc.generate_dashboard(pd.DataFrame(self.MARGIN_SHEET), 2026, 'all', 1, 12)
        built = {c['chartType'] for c in r['charts']}
        assert built == {'gross_margin_trend'}, f'毛利率百分比被当成金额进了这些图: {built - {"gross_margin_trend"}}'

    def test_expense_sheet_name_marks_rows_as_expense(self):
        r = svc.generate_dashboard(pd.DataFrame(self.EXPENSE_SHEET), 2026, 'all', 1, 12, sheet_name='费用预算')
        charts = {c['chartType']: c for c in r['charts']}
        assert charts['expense_yoy_budget']['success'] is True
        assert to_yuan(kpi(charts['expense_yoy_budget'], '本年费用合计')['value']) == 2_700_000
        # 费用表不能出现在「业绩」预算完成 / 同比里
        assert charts['budget_achievement']['success'] is False
        assert charts['yoy_mom_comparison']['success'] is False

    def test_margin_without_cost_line_is_not_100_percent(self):
        rows = [{'月份': f'{m}月', '本年（万元）': 100 * m, '上年（万元）': 80 * m} for m in range(1, 4)]
        r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12, sheet_name='业绩同比')
        card = next(c for c in r['charts'] if c['chartType'] == 'kpi_scorecard')
        labels = {k['label'] for k in card['kpis']}
        assert not labels & {'毛利率', '净利率'}, f'没有成本科目却给出了利润率: {card["kpis"]}'

    def test_sankey_and_waterfall_agree_on_net_margin(self):
        rows = [{'月份': '1月', '项目': k, '实际': v} for k, v in PNL_REFERENCE]
        r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)
        charts = {c['chartType']: c for c in r['charts']}
        waterfall_margin = charts['pnl_waterfall']['metadata']['netMargin']
        sankey_margin = charts['cost_flow_sankey']['metadata']['netMargin']
        assert waterfall_margin == pytest.approx(3308 / 8333 * 100, abs=0.01)
        assert sankey_margin == pytest.approx(waterfall_margin, abs=0.01)

    def test_scorecard_counts_annual_rows_without_months(self):
        rows = [{'品类': '酱料', '本年（万元）': 30, '上年（万元）': 20}, {'品类': '汤料', '本年（万元）': 70, '上年（万元）': 60}]
        r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)
        card = next(c for c in r['charts'] if c['chartType'] == 'kpi_scorecard')
        assert to_yuan(kpi(card, '实际总额')['value']) == 1_000_000

    def test_scorecard_on_plain_sheet_shows_actual_total_not_zero_revenue(self):
        r = svc.generate_dashboard(pd.DataFrame(self.EXPENSE_SHEET), 2026, 'all', 1, 12)
        card = next(c for c in r['charts'] if c['chartType'] == 'kpi_scorecard')
        labels = [k['label'] for k in card['kpis']]
        assert '营业收入' not in labels, card['kpis']
        assert to_yuan(kpi(card, '实际总额')['value']) == 2_700_000


class TestDisplayConsistency:
    """设计评审 (2026-10-21) 发现: 同一个数在两张图里一张显示元、一张显示万; 月份筛选排成 10月…9月。"""

    def test_same_total_has_same_unit_across_charts(self):
        rows = [{'月份': f'{m}月', '预算': 7000, '实际': 7050, '上年': 6000} for m in range(1, 7)]
        r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)
        charts = {c['chartType']: c for c in r['charts']}
        budget_actual = kpi(charts['budget_achievement'], '年度实际')['value']
        variance_actual = kpi(charts['variance_analysis'], '总实际')['value']
        assert budget_actual == variance_actual == '4.23万', (budget_actual, variance_actual)

    def test_month_filter_values_are_in_calendar_order(self, monkeypatch):
        client, _ = _tenant_client(monkeypatch, 'F006')
        body = client.post('/batch', json={'upload_id': 5247, 'period_type': 'all'}).json()
        months = next(d['values'] for d in body['availableDimensions'] if d['name'] == '月份')
        assert months == [f'{m}月' for m in range(1, 13)]


class TestHeadline:
    """页首一句话结论: 规则生成 (不靠 AI), 口径与图表同一份指标库。"""

    BUDGET = [{'月份': f'{m}月', '预算（万元）': t, '实际（万元）': a} for m, t, a in [
        (1, 1000, 200), (2, 1000, 300), (3, 1000, 400), (4, 1000, 1100), (5, 1000, 900), (6, 1000, 950)]]

    def test_budget_sheet_states_achievement_and_weakest_quarter(self):
        r = svc.generate_dashboard(pd.DataFrame(self.BUDGET), 2026, 'all', 1, 12)
        h = r['headline']
        assert h['tone'] == 'bad'
        assert '实际 3,850.00万' in h['text'] and '目标 6,000.00万' in h['text']
        assert '达成 64.2%' in h['text'] and '未达' in h['text']
        assert 'Q1 30.0%' in h['text'] and '4月 110.0%' in h['text']

    def test_pnl_sheet_states_revenue_and_net_profit(self):
        rows = [{'月份': '1月', '项目': k, '实际（万元）': v} for k, v in PNL_REFERENCE]
        h = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['headline']
        assert '营业收入 8,333.00万' in h['text'] and '净利润 3,308.00万' in h['text'] and '净利率 39.7%' in h['text']

    def test_expense_sheet_calls_overrun_an_overrun(self):
        rows = [{'月份': f'{m}月', '预算': 100, '实际': 120, '上年': 110} for m in range(1, 4)]
        h = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12, sheet_name='费用预算')['headline']
        assert '费用' in h['text'] and '超支' in h['text'] and h['tone'] == 'bad'

    def test_ratio_sheet_has_no_amount_headline(self):
        rows = [{'月份': '1月', '本年毛利率': '25%', '上年毛利率': '20%'}]
        h = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['headline']
        assert h is None or '万' not in h['text']


BUDGET_12 = [{'月份': f'{m}月', '预算（万元）': 1000, '实际（万元）': 100 * m} for m in range(1, 13)]
GENERIC_GAP = '这份数据缺少这张图需要的列'


class TestMissingChartsExplained:
    """设计复评 (2026-10-23): 「还缺数据」12 条里 9 条是同一句通用话, 还一律叫人重新上传。
    出不了的图要写明缺哪一列, 并分清「补一列」和「需要另一张表」。"""

    def _dashboard(self, rows, **kw):
        return svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12, **kw)

    def test_unbuildable_charts_name_the_missing_columns(self):
        r = self._dashboard(BUDGET_12)
        gaps = {u['chartType']: u for u in r['unavailableCharts']}
        built = {c['chartType'] for c in r['charts']}
        assert 'budget_achievement' in built and 'budget_achievement' not in gaps
        yoy = gaps['yoy_mom_comparison']
        assert yoy['gapKind'] == 'column' and '「上年」' in yoy['reason'], yoy
        cat = gaps['category_yoy_comparison']
        assert '「上年」' in cat['reason'] and '「品类' in cat['reason'], cat
        assert all(GENERIC_GAP not in u['reason'] for u in gaps.values())

    def test_charts_that_need_another_statement_say_which(self):
        gaps = {u['chartType']: u for u in self._dashboard(BUDGET_12)['unavailableCharts']}
        pnl = gaps['pnl_waterfall']
        assert pnl['gapKind'] == 'statement' and '损益表' in pnl['reason'] and '营业收入' in pnl['reason'], pnl
        assert gaps['cash_flow_waterfall']['gapKind'] == 'statement'
        assert '现金流量表' in gaps['cash_flow_waterfall']['reason']

    def test_built_but_empty_charts_carry_a_gap_kind(self):
        charts = {c['chartType']: c for c in self._dashboard(BUDGET_12)['charts']}
        assert charts['channel_analysis']['metadata']['gapKind'] == 'column'
        assert '「渠道」' in charts['channel_analysis']['metadata']['reason']
        assert charts['hr_cost_analysis']['metadata']['gapKind'] == 'statement'
        # 出了图的不带 gapKind (阴性对照)
        assert 'gapKind' not in charts['budget_achievement']['metadata']

    def test_ratio_sheet_explains_why_amount_charts_are_missing(self):
        rows = [{'月份': f'{m}月', '本年毛利率': '25%', '上年毛利率': '20%'} for m in range(1, 4)]
        gaps = {u['chartType']: u for u in self._dashboard(rows)['unavailableCharts']}
        assert gaps['budget_achievement']['gapKind'] == 'statement'
        assert '比率' in gaps['budget_achievement']['reason']
        assert 'gross_margin_trend' not in gaps


class TestBudgetChartIsReadable:
    """设计复评 (2026-10-23): 预算达成图叠了柱、两条达成率线、季度底色+百分比、季度虚线均值、
    一排彩色圆点、目标线和缩放条。减到柱 + 一条达成率线 + 100% 目标线; 季度达成率在上方季度条里。"""

    def _chart(self):
        r = svc.generate_dashboard(pd.DataFrame(BUDGET_12), 2026, 'all', 1, 12)
        return next(c for c in r['charts'] if c['chartType'] == 'budget_achievement')

    def test_only_bars_one_rate_line_and_target_line(self):
        opt = self._chart()['echartsOption']
        names = [s['name'] for s in opt['series']]
        assert len(names) == 3 and names[2] == '达成率', names
        assert '累计达成率' not in names and '累计达成率' not in opt['legend']['data']
        assert 'graphic' not in opt and 'dataZoom' not in opt
        assert all(isinstance(s['itemStyle']['color'], str) for s in opt['series'][:2])
        assert 'markLine' not in opt['series'][1]
        target = opt['series'][2]['markLine']['data'][0]
        assert target['yAxis'] == 100

    def test_quarter_bands_have_no_duplicate_rates_and_numbers_are_kept(self):
        chart = self._chart()
        bands = chart['echartsOption']['series'][0].get('markArea', {}).get('data', [])
        assert bands and all('%' not in str(b[0].get('label', {}).get('formatter', '')) for b in bands)
        assert [q['rate'] for q in chart['quarterlyProgress']] == [20.0, 50.0, 80.0, 110.0]
        rows = {row['label'] for row in chart['tableData']['rows']}
        assert '累计达成率(%)' in rows


def _charts(rows, **kw):
    r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12, **kw)
    return r, {c['chartType']: c for c in r['charts']}


class TestMonthOverMonth:
    """客户需求「实际与同期、环比」: 环比原来只在图顶部一排没标签的小字里 (环比线因默认「综合」模式永远不画)。"""

    ROWS = [{'月份': f'{m}月', '实际': a, '上年': ly} for m, a, ly in [(1, 100, 90), (2, 120, 100), (3, 90, 95), (4, 108, 100)]]

    def test_mom_line_is_drawn_and_named(self):
        opt = _charts(self.ROWS)[1]['yoy_mom_comparison']['echartsOption']
        mom = next(s for s in opt['series'] if s['name'] == '环比增长率')
        assert mom['type'] == 'line'
        assert mom['data'][:4] == [None, 20.0, -25.0, 20.0]
        assert all(v is None for v in mom['data'][4:])  # 没上传的月份是空, 不是 0
        assert '环比增长率' in opt['legend']['data']

    def test_latest_mom_is_a_kpi_with_its_month(self):
        kpis = {k['label']: k for k in _charts(self.ROWS)[1]['yoy_mom_comparison']['kpis']}
        assert kpis['最近月环比']['value'] == '4月 +20.0'
        assert kpis['最近月环比']['unit'] == '%'

    def test_unlabeled_mom_row_is_gone(self):
        opt = _charts(self.ROWS)[1]['yoy_mom_comparison']['echartsOption']
        top_row = [g for g in opt.get('graphic', []) if g.get('top') == '7%']
        assert top_row == []

    def test_single_month_has_no_mom_kpi(self):
        kpis = {k['label']: k for k in _charts(self.ROWS[:1])[1]['yoy_mom_comparison']['kpis']}
        assert kpis['最近月环比']['value'] == '-'


class TestBusinessUnits:
    """客户需求「从各业务单元」: 能识别事业部 / 部门列, 但原来没有对应的图。"""

    ROWS = [
        {'月份': f'{m}月', '事业部': bu, '预算': b, '实际': a, '上年': ly}
        for m in (1, 2)
        for bu, b, a, ly in [('华东', 500, 600, 450), ('华南', 400, 300, 320), ('西南', 100, 100, 50)]
    ]

    def test_units_are_compared_by_actual_share_achievement_and_yoy(self):
        chart = _charts(self.ROWS)[1]['business_unit_analysis']
        kpis = {k['label']: k['value'] for k in chart['kpis']}
        assert kpis['业务单元数'] == '3个'
        assert kpis['实际最大'] == '华东 (60.0%)'
        assert kpis['达成最低'] == '华南 (75.0%)'
        assert kpis['同比增长最快'] == '西南 (+100.0%)'
        opt = chart['echartsOption']
        assert opt['yAxis']['data'] == ['西南', '华南', '华东']  # 横向条形, 最大的在最上面
        assert [s['name'] for s in opt['series']][:2] == ['实际', '预算']

    def test_pnl_sheet_counts_revenue_rows_only(self):
        rows = [{'月份': '1月', '事业部': bu, '项目': item, '实际': v}
                for bu, item, v in [('华东', '营业收入', 1000), ('华东', '营业成本', 600), ('华南', '营业收入', 500), ('华南', '营业成本', 450)]]
        chart = _charts(rows)[1]['business_unit_analysis']
        kpis = {k['label']: k['value'] for k in chart['kpis']}
        assert kpis['实际最大'] == '华东 (66.7%)'

    def test_missing_unit_column_is_explained(self):
        r, charts = _charts(BUDGET_12)
        assert 'business_unit_analysis' not in charts
        gap = next(u for u in r['unavailableCharts'] if u['chartType'] == 'business_unit_analysis')
        assert gap['gapKind'] == 'column' and '「事业部' in gap['reason']


class TestFindings:
    """客户需求「好的与不好的、建议」: 原来只有 AI 生成; 现在先给规则算出来的, 每条都带数字。"""

    def test_budget_sheet_names_good_quarters_and_the_weak_one(self):
        f = svc.generate_dashboard(pd.DataFrame(BUDGET_12), 2026, 'all', 1, 12)['findings']
        assert any('Q4 达成 110.0%' in g for g in f['good'])
        assert any('Q1 只达成 20.0%' in a for a in f['attention'])
        assert any(s.startswith('Q1') and '缺口' in s for s in f['suggestions'])
        for line in f['good'] + f['attention'] + f['suggestions']:
            assert re.search(r'\d', line), line

    def test_best_month_is_good_news_even_when_no_quarter_reaches_target(self):
        # 线上 5250 (2026-10-24): 四个季度都没达标, 但 8 月 124.4%, 「做得好的」原来是空的
        actuals = [100, 200, 300, 400, 500, 600, 700, 1244, 900, 500, 500, 500]
        rows = [{'月份': f'{m}月', '预算': 1000, '实际': a} for m, a in zip(range(1, 13), actuals)]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert not any(g.startswith('Q') for g in f['good'])
        assert '8月 达成 124.4%，各月里最好' in f['good']

    def test_no_month_reaches_target_means_no_month_good_news(self):
        rows = [{'月份': f'{m}月', '预算': 1000, '实际': 500} for m in range(1, 13)]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert not any('各月里最好' in g for g in f['good'])

    def test_expense_sheet_flags_the_overrun_months(self):
        rows = [{'月份': f'{m}月', '预算': 100, '实际': a, '上年': 110} for m, a in [(1, 90), (2, 280), (3, 120)]]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12, sheet_name='费用预算')['findings']
        assert any('超预算' in a and '163.3%' in a for a in f['attention'])
        assert any('2月' in s and '280.0%' in s for s in f['suggestions'])
        assert not any('超支' in g for g in f['good'])

    def test_yoy_decline_months_are_attention(self):
        f = svc.generate_dashboard(pd.DataFrame(TestMonthOverMonth.ROWS), 2026, 'all', 1, 12)['findings']
        assert any('同比增长' in g for g in f['good'])
        assert any('3月' in a and '比上年' in a for a in f['attention'])

    def test_business_unit_below_target_gets_a_suggestion(self):
        f = svc.generate_dashboard(pd.DataFrame(TestBusinessUnits.ROWS), 2026, 'all', 1, 12)['findings']
        assert any('华南' in a and '75.0%' in a for a in f['attention'])
        assert any('华南' in s for s in f['suggestions'])

    def test_ratio_sheet_has_no_amount_findings(self):
        rows = [{'月份': '1月', '本年毛利率': '25%', '上年毛利率': '20%'}]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert f is None

    # 样图⑤「各月毛利率」(比率表) 原来没有规则结论, 结论区只能靠 AI。
    MARGIN_CUR = [24, 25, 26, 27, 28, 25, 25, 27, 29, 28, 24, 25]
    MARGIN_LY = [25, 24, 25, 26, 27, 27, 26, 26, 27, 27, 27, 24]
    MARGIN_ROWS = [{'月份': f'{m}月', '本年毛利率': f'{c}%', '上年毛利率': f'{ly}%'}
                   for m, c, ly in zip(range(1, 13), MARGIN_CUR, MARGIN_LY)]

    def test_margin_sheet_names_average_best_worst_and_declines(self):
        r = svc.generate_dashboard(pd.DataFrame(self.MARGIN_ROWS), 2026, 'all', 1, 12)
        f = r['findings']
        assert '各月平均毛利率 26.1%，比上年 +0.2 个百分点（各月简单平均）' in f['good']
        assert '9月毛利率最高，29.0%' in f['good']
        assert '1月毛利率最低，24.0%，比最高月低 5.0 个百分点' in f['attention']
        declines = [a for a in f['attention'] if '比上年低' in a]
        assert declines == ['11月比上年低 3.0 个百分点（24.0%，上年 27.0%）',
                            '6月比上年低 2.0 个百分点（25.0%，上年 27.0%）']
        assert f['suggestions'] and f['suggestions'][0].startswith('复盘 11月')
        for line in f['good'] + f['attention'] + f['suggestions']:
            assert '万' not in line, line  # 比率表不出金额结论
        # 与毛利率趋势图同一口径: 月均毛利率 KPI 也是 26.1
        kpis = {k['label']: k['value'] for k in _charts(self.MARGIN_ROWS)[1]['gross_margin_trend']['kpis']}
        assert kpis['本年月均毛利率'] == '26.1'

    def test_margin_sheet_headline_states_the_average_gap(self):
        h = svc.generate_dashboard(pd.DataFrame(self.MARGIN_ROWS), 2026, 'all', 1, 12)['headline']
        assert '各月平均毛利率 26.1%，上年 25.9%，+0.2 个百分点（各月简单平均）' in h['text']
        assert h['tone'] == 'good'
        worse = [{**row, '本年毛利率': '20%'} for row in self.MARGIN_ROWS]
        assert svc.generate_dashboard(pd.DataFrame(worse), 2026, 'all', 1, 12)['headline']['tone'] == 'bad'

    def test_margin_gap_is_computed_from_the_displayed_averages(self):
        # 线上 5254 (2026-10-25): 「各月平均毛利率 26.3%，上年 26.3%，+0.1 个百分点」读起来对不上
        rows = [{'月份': f'{m}月', '本年毛利率': '26.34%', '上年毛利率': '26.26%'} for m in (1, 2)]
        # 改成「与上年持平」(2026-10-25): 「+0.0 个百分点」放在「做得好的」里读起来像在表扬持平
        r = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)
        assert '各月平均毛利率 26.3%，与上年持平（各月简单平均）' in r['headline']['text']
        assert '+0.0' not in r['headline']['text']
        assert r['headline']['tone'] == 'neutral'
        f = r['findings']
        assert not any('平均毛利率' in line for line in f['good'] + f['attention'] + f['suggestions'])
        # 阳性对照: 差 0.1 个百分点时照常给出差值并算好消息
        up = [{**row, '本年毛利率': '26.44%'} for row in rows]
        r = svc.generate_dashboard(pd.DataFrame(up), 2026, 'all', 1, 12)
        assert '各月平均毛利率 26.4%，上年 26.3%，+0.1 个百分点' in r['headline']['text']
        assert r['headline']['tone'] == 'good'
        assert '各月平均毛利率 26.4%，比上年 +0.1 个百分点（各月简单平均）' in r['findings']['good']

    def test_pnl_sheet_states_margin_expense_share_and_biggest_item(self):
        # 线上 5252 (2026-10-25): 损益表只有「净利润」一条结论
        rows = [{'月份': '1月', '项目': k, '实际（万元）': v} for k, v in PNL_REFERENCE]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert '净利润 3,308.00万，净利率 39.7%' in f['good']
        assert '毛利率 82.3%（毛利 6,862.00万 / 营业收入 8,333.00万）' in f['good']
        assert ('期间费用 2,472.00万，占营业收入 29.7%，最大一项是研发费用 999.00万（占 40%）'
                in f['attention'])
        assert '营业外收支净减少 44.00万（收入 42.00万，支出 86.00万）' in f['attention']
        assert f['suggestions'] == ['先看研发费用：999.00万，占期间费用 40%，是期间费用里最大的一项']

    def test_pnl_with_last_year_states_net_profit_growth(self):
        rows = [{'月份': '1月', '项目': k, '实际（万元）': v, '上年（万元）': v * 0.8} for k, v in PNL_REFERENCE]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert any(g.startswith('净利润同比 +25.0%') for g in f['good'])

    def test_yoy_sheet_names_best_month_and_suggests_reviewing_the_biggest_drop(self):
        # 线上 5251 (2026-10-25): 同比表没有「建议」
        rows = [{'月份': f'{m}月', '实际': a, '上年': ly}
                for m, a, ly in [(1, 100, 90), (2, 150, 100), (3, 80, 120), (4, 95, 100)]]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert '2月同比 +50.0%（多 50.00），各月里增长最多' in f['good']
        assert f['suggestions'] == ['复盘 3月：比上年少 40.00，各月里减少的金额最多']

    def test_latest_mom_finding_matches_the_kpi_month(self):
        f = svc.generate_dashboard(pd.DataFrame(TestMonthOverMonth.ROWS), 2026, 'all', 1, 12)['findings']
        assert '4月环比 +20.0%（上月 90.00 → 本月 108.00）' in f['good']  # 列名没写单位, 金额不带「万」
        falling = TestMonthOverMonth.ROWS[:3]  # 最近一月是 3 月: 120 → 90
        f = svc.generate_dashboard(pd.DataFrame(falling), 2026, 'all', 1, 12)['findings']
        assert any(a.startswith('3月环比 -25.0%') for a in f['attention'])
        assert not any('环比' in g for g in f['good'])


    # 样图⑥ 品类结构: 原来只有「占比变化」两条, 没有建议 (线上 5255, 2026-10-25)
    CATEGORY_ROWS = [{'品类': k, '月份': '1月', '本年': c, '上年': ly}
                     for k, c, ly in [('酱料', 600, 400), ('汤料', 250, 300), ('底料', 150, 100)]]

    def test_category_sheet_names_growth_decline_concentration_and_suggests_review(self):
        f = svc.generate_dashboard(pd.DataFrame(self.CATEGORY_ROWS), 2026, 'all', 1, 12)['findings']
        assert '酱料比上年多 200.00（+50.0%），各品类里增长最多' in f['good']
        assert '汤料比上年少 50.00（-16.7%）' in f['attention']
        assert '酱料占本期 60.0%，一半以上来自这一个品类' in f['attention']
        assert '复盘汤料：比上年少 50.00，各品类里减少的金额最多' in f['suggestions']

    def test_category_without_decline_suggests_the_share_loser_not_a_review(self):
        rows = [{'品类': k, '月份': '1月', '本年': c, '上年': ly}
                for k, c, ly in [('酱料', 300, 200), ('汤料', 220, 200), ('底料', 280, 200)]]
        f = svc.generate_dashboard(pd.DataFrame(rows), 2026, 'all', 1, 12)['findings']
        assert not any('复盘' in s and '品类' in s for s in f['suggestions'])
        assert '关注汤料：占比比上年低 5.8 个百分点，金额同比 +10.0%，增长慢于整体' in f['suggestions']
        assert not any('一半以上' in a for a in f['attention'])


class TestSlicerScope:
    """线上验收 (2026-10-22): 点筛选器「3月」后数字已是 3 月的, 页首结论和期间标签仍写「全部历史」。
    期间标签要说清楚这些数字是哪一部分数据。走真实 /batch 入口 (筛选在接口层生效)。"""

    def test_month_filter_is_named_in_headline_and_period_label(self, monkeypatch):
        client, _ = _tenant_client(monkeypatch, 'F006')
        body = client.post('/batch', json={'upload_id': 5245, 'period_type': 'all',
                                           'filters': {'月份': ['2月']}}).json()
        assert body['period']['label'] == '全部历史 · 月份 2月', body['period']
        text = body['headline']['text']
        assert text.startswith('全部历史 · 月份 2月：'), text
        assert '实际 240.00' in text and '目标 220.00' in text, text

    def test_no_filter_keeps_plain_label(self, monkeypatch):
        client, _ = _tenant_client(monkeypatch, 'F006')
        body = client.post('/batch', json={'upload_id': 5245, 'period_type': 'all'}).json()
        assert body['period']['label'] == '全部历史'
        assert body['headline']['text'].startswith('全部历史：')
        assert '实际 420.00' in body['headline']['text']

    def test_filter_on_a_column_the_data_lacks_is_not_claimed(self, monkeypatch):
        # 接口层忽略不存在的列 (数据没被筛), 标签就不能说筛过
        client, _ = _tenant_client(monkeypatch, 'F006')
        body = client.post('/batch', json={'upload_id': 5245, 'period_type': 'all',
                                           'filters': {'渠道': ['线上'], '月份': []}}).json()
        assert body['period']['label'] == '全部历史'


class TestExpenseChart:
    def test_expense_total_excludes_revenue(self):
        r = build('expense_yoy_budget', [
            {'月份': '1月', '项目': '营业收入', '预算': 1000, '实际': 1000, '上年': 900},
            {'月份': '1月', '项目': '销售费用', '预算': 100, '实际': 100, '上年': 90},
        ])
        assert to_yuan(kpi(r, '本年费用合计')['value']) == 100

    def test_rate_axis_fits_large_overruns(self):
        r = build('expense_yoy_budget', [
            {'月份': '1月', '项目': '销售费用', '预算': 100, '实际': 277.73, '上年': 90},
        ])
        axes = r['echartsOption']['yAxis']
        rate_axis = next(a for a in axes if '%' in str(a.get('name', '')))
        assert rate_axis.get('max') is None or rate_axis['max'] >= 277.73


class TestMissingVersusZero:
    def test_missing_month_is_no_data_not_minus_100_percent(self):
        r = build('yoy_mom_comparison', [
            {'月份': '1月', '实际': 100, '上年': 0},
            {'月份': '3月', '实际': 200, '上年': 100},
        ], 1, 3)
        assert table_row(r, '本年')[1] is None, '没上传的 2 月不能显示成 0'
        assert table_row(r, '环比(%)') == [None, None, None]
        # 上年 1 月是真实的 0: 增长率无法计算, 不能显示 0%
        assert table_row(r, '同比(%)')[0] is None


class TestSuccessReflectsChartState:
    def test_caught_builder_error_is_not_counted_as_success(self):
        from smartbi.services.financial_data_normalizer import ColumnMapping
        # 有渠道列但缺「实际」列: 构建器内部报错并自己捕获, 返回 dataQuality=error
        r = registry._safe_build(registry.get('channel_analysis'), pd.DataFrame({'channel': ['商超']}),
                                 ColumnMapping(channel_col='渠道'), {'start_month': 1, 'end_month': 1}, 2026)
        assert r['metadata']['dataQuality'] == 'error'
        assert r['success'] is False

    def test_dashboard_success_count_excludes_insufficient_charts(self):
        r = svc.generate_dashboard(pd.DataFrame([{'月份': '1月', '品类': '酱料', '实际': 100}]),
                                   2026, 'month_range', 1, 1)
        ok = [c for c in r['charts'] if c.get('success')]
        assert all(c.get('metadata', {}).get('dataQuality') not in ('error', 'insufficient') for c in ok)
        assert r['successCount'] == len(ok)


# ---------------------------------------------------------------------------
# 第三轮
# ---------------------------------------------------------------------------
class TestInputParsing:
    def test_unit_in_column_header_is_applied(self):
        r = build('budget_achievement', [{'月份': '1月', '预算（万元）': 20, '实际（万元）': 10}])
        assert to_yuan(kpi(r, '年度实际')['value']) == 100000
        assert to_yuan(kpi(r, '年度目标')['value']) == 200000

    def test_iso_day_dates_map_to_their_month(self):
        r = build('budget_achievement', [{'日期': '2026-01-15', '预算': 200, '实际': 100}])
        assert to_yuan(kpi(r, '年度实际')['value']) == 100

    def test_percentage_strings_are_numbers(self):
        r = build('gross_margin_trend', [{'月份': '1月', '本年毛利率': '25%', '上年毛利率': '20%'}])
        assert r['metadata']['dataQuality'] != 'insufficient', r['analysisContext']
        assert float(kpi(r, '平均差异')['value']) == pytest.approx(5.0)


class TestProductRanking:
    ROWS = [
        {'月份': '1月', '品类': '酱料', '产品名称': '产品A', '实际': 100},
        {'月份': '1月', '品类': '酱料', '产品名称': '产品B', '实际': 200},
        {'月份': '1月', '品类': '退货品类', '产品名称': '产品C', '实际': -50},
    ]

    def test_ranks_individual_products_not_categories(self):
        r = build('product_ranking', self.ROWS)
        assert kpi(r, 'Top 1 产品')['value'] == '产品B'
        assert kpi(r, '产品总数')['value'] == '3'

    def test_negative_revenue_is_kept_and_net_total_is_used(self):
        r = build('product_ranking', self.ROWS)
        rows = r['tableData']['rows']
        names = [row['label'] if isinstance(row, dict) else row[1] for row in rows]
        assert '产品C' in names, '退货产品不能从排名里消失'
        assert r['metadata']['netRevenue'] == pytest.approx(250)


class TestNoSynthesizedCharts:
    """同因扫描: 缺数据时按预设比例拆分 / 随机波动「造」出来的图, 一律改为明确提示缺什么。"""

    SALES_ONLY = [
        {'月份': f'{m}月', '预算': 100 * m, '实际': 90 * m, '上年': 80 * m} for m in range(1, 7)
    ]
    SYNTH_PRONE = ('ar_aging', 'cashflow_trend', 'cash_flow_waterfall', 'hr_cost_analysis')

    def test_sales_sheet_does_not_produce_invented_cash_ar_or_hr_charts(self):
        r = svc.generate_dashboard(pd.DataFrame(self.SALES_ONLY), 2026, 'month_range', 1, 6)
        for c in r['charts']:
            if c['chartType'] in self.SYNTH_PRONE:
                assert c['success'] is False, f"{c['chartType']} 在没有对应数据时生成了图表: {c.get('kpis')}"
                assert c['metadata'].get('dataQuality') in ('insufficient', 'error')
                assert c.get('analysisContext'), '必须告诉用户缺什么'

    def test_same_input_gives_same_numbers(self):
        a = svc.generate_dashboard(pd.DataFrame(self.SALES_ONLY), 2026, 'month_range', 1, 6)
        b = svc.generate_dashboard(pd.DataFrame(self.SALES_ONLY), 2026, 'month_range', 1, 6)
        ka = {c['chartType']: c.get('kpis') for c in a['charts']}
        kb = {c['chartType']: c.get('kpis') for c in b['charts']}
        assert ka == kb, '同一份数据两次生成的数字不同 (含随机数)'

    def test_real_aging_columns_are_used(self):
        r = build('ar_aging', [
            {'客户': '甲', '实际': 600, '0-30天': 300, '31-60天': 200, '61-90天': 100},
        ])
        assert r['success'] is True
        assert r['metadata']['totalAR'] == pytest.approx(600)

    def test_real_cash_flow_items_are_used(self):
        r = build('cashflow_trend', [
            {'月份': '1月', '项目': '经营活动现金流量净额', '实际': 500},
            {'月份': '1月', '项目': '投资活动现金流量净额', '实际': -200},
            {'月份': '1月', '项目': '筹资活动现金流量净额', '实际': 50},
        ])
        assert r['success'] is True
        assert r['metadata']['dataQuality'] == 'good'
        headers = r['tableData']['headers']
        net_col = next(i for i, h in enumerate(headers) if str(h).startswith('净'))
        january = r['tableData']['rows'][0]
        assert to_yuan(january[net_col]) == pytest.approx(350)


# ---------------------------------------------------------------------------
# 导出
# ---------------------------------------------------------------------------
def _export_client():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from smartbi.api import financial_dashboard as api

    app = FastAPI()
    @app.middleware('http')
    async def authenticated_factory(request,call_next):
        request.state.factory_id='F006'
        request.state.role='factory_super_admin'
        return await call_next(request)
    app.include_router(api.router)
    return TestClient(app)


def _sheet_cells(content: bytes) -> List[List[Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content))
    cells = []
    for ws in wb:
        cells.extend([list(r) for r in ws.values])
    return cells


class TestExcelExport:
    def test_scaled_table_states_its_unit(self):
        chart = build('budget_achievement', [{'月份': '1月', '预算': 20000, '实际': 10000}])
        resp = _export_client().post('/export-excel', json={'charts': [chart]})
        assert resp.status_code == 200
        flat = ' '.join(str(c) for row in _sheet_cells(resp.content) for c in row if c is not None)
        assert '万元' in flat, '按万显示的明细必须写明单位'

    def test_product_ranking_table_exports(self):
        chart = build('product_ranking', [
            {'月份': '1月', '品类': '酱料', '产品名称': '产品A', '实际': 100},
            {'月份': '1月', '品类': '汤料', '产品名称': '产品B', '实际': 200},
        ])
        resp = _export_client().post('/export-excel', json={'charts': [chart]})
        assert resp.status_code == 200
        assert 'spreadsheetml' in resp.headers.get('content-type', '')
        flat = [c for row in _sheet_cells(resp.content) for c in row]
        assert '产品A' in flat

    def test_report_carries_the_headline_and_rule_findings_first(self):
        # 导出报告原来只有 KPI / 明细 / AI 分析, 页面上的页首结论和规则结论带不出去
        r = svc.generate_dashboard(pd.DataFrame(TestFindings.CATEGORY_ROWS), 2026, 'all', 1, 12)
        resp = _export_client().post('/export-excel', json={
            'charts': r['charts'], 'headline': r['headline'], 'findings': r['findings']})
        assert resp.status_code == 200
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        assert wb.sheetnames[0] == '分析结论'
        rows = [list(x) for x in wb['分析结论'].values][1:]
        assert rows[0] == ['页首结论', r['headline']['text']]
        expected = ([['做得好的', g] for g in r['findings']['good']]
                    + [['需要关注', a] for a in r['findings']['attention']]
                    + [['建议', s] for s in r['findings']['suggestions']])
        assert rows[1:] == expected
        assert ['建议', '复盘汤料：比上年少 50.00，各品类里减少的金额最多'] in rows

    def test_no_conclusions_means_no_conclusion_sheet(self):
        chart = build('budget_achievement', [{'月份': '1月', '预算': 200, '实际': 100}])
        resp = _export_client().post('/export-excel', json={'charts': [chart], 'findings': None})
        import openpyxl
        assert openpyxl.load_workbook(io.BytesIO(resp.content)).sheetnames[0] == 'KPI汇总'

    def test_export_failure_is_not_http_200(self, monkeypatch):
        import openpyxl

        def boom(self, *a, **k):
            raise RuntimeError('disk full')

        monkeypatch.setattr(openpyxl.Workbook, 'save', boom)
        resp = _export_client().post('/export-excel', json={'charts': []})
        assert resp.status_code >= 500
        assert resp.json()['success'] is False


# ---------------------------------------------------------------------------
# 上传读取 + 租户
# ---------------------------------------------------------------------------
UPLOADED_ROWS = {
    # (upload_id, factory_id) -> (sheet_name, rows as stored by the uploader: one row_data per Excel row)
    (5245, 'F006'): ('预算完成', [
        {'月份': '1月', '预算': 200, '实际': 180, '上年': 150},
        {'月份': '2月', '预算': 220, '实际': 240, '上年': 160},
    ]),
    (5247, 'F006'): ('全年业绩', [
        {'月份': f'{m}月', '项目': f'品类{m % 3}', '实际': 100 * m} for m in range(1, 13)
    ]),
    (5246, 'F006'): ('费用预算', [
        {'月份': '1月', '预算': 100, '实际': 90, '上年': 80},
        {'月份': '2月', '预算': 100, '实际': 95, '上年': 85},
    ]),
}


class _FakeConn:
    def __init__(self, calls):
        self.calls = calls

    async def fetch(self, sql, *args):
        self.calls.append((sql, args))
        assert 'row_data' in sql and 'upload_id' in sql, sql
        upload_id, factory_id = args[0], args[1]
        sheet, rows = UPLOADED_ROWS.get((upload_id, factory_id), ('', []))
        return [{'row_data': r, 'sheet_name': sheet} for r in rows]


class _FakePool:
    def __init__(self):
        self.calls = []

    def acquire(self):
        pool = self

        class _Ctx:
            async def __aenter__(self):
                return _FakeConn(pool.calls)

            async def __aexit__(self, *exc):
                return False
        return _Ctx()


def _tenant_client(monkeypatch, token_factory: Optional[str]):
    from fastapi import FastAPI, Request
    from fastapi.testclient import TestClient
    from smartbi.api import financial_dashboard as api

    pool = _FakePool()

    async def fake_get_pool():
        return pool

    monkeypatch.setattr('smartbi.config.get_pg_pool', fake_get_pool)
    app = FastAPI()

    @app.middleware('http')
    async def inject_jwt_claims(request: Request, call_next):
        # 模拟 JWTAuthMiddleware: factory_id 来自登录身份
        request.state.factory_id = token_factory
        request.state.auth_method = 'jwt'
        return await call_next(request)

    app.include_router(api.router)
    return TestClient(app), pool


class TestUploadLoading:
    def test_uploaded_rows_are_read_by_upload_id_for_token_factory(self, monkeypatch):
        client, pool = _tenant_client(monkeypatch, 'F006')
        resp = client.post('/batch', json={'upload_id': 5245, 'year': 2026, 'period_type': 'month_range',
                                           'start_month': 1, 'end_month': 2})
        body = resp.json()
        assert resp.status_code == 200, body
        assert body['success'] is True
        budget = next(c for c in body['charts'] if c['chartType'] == 'budget_achievement')
        assert to_yuan(kpi(budget, '年度实际')['value']) == 420
        assert pool.calls and pool.calls[0][1][:2] == (5245, 'F006')

    def test_uploaded_sheet_name_reaches_the_charts(self, monkeypatch):
        # 「费用预算」表没有科目列: 靠库里记录的工作表名才知道它是费用, 不能出现在业绩图里
        client, _ = _tenant_client(monkeypatch, 'F006')
        body = client.post('/batch', json={'upload_id': 5246, 'period_type': 'all'}).json()
        charts = {c['chartType']: c for c in body['charts']}
        assert charts['expense_yoy_budget']['success'] is True
        assert charts['budget_achievement']['success'] is False

    def test_other_factory_cannot_read_the_upload(self, monkeypatch):
        client, _ = _tenant_client(monkeypatch, 'F001')
        resp = client.post('/batch', json={'upload_id': 5245})
        assert resp.status_code == 404
        assert resp.json()['success'] is False

    def test_body_factory_id_cannot_override_token(self, monkeypatch):
        client, pool = _tenant_client(monkeypatch, 'F001')
        resp = client.post('/batch', json={'upload_id': 5245, 'factory_id': 'F006'})
        assert resp.status_code == 403
        assert not pool.calls, '越权请求不应触发查询'
