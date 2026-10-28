"""Financial Dashboard API — Chart generation, batch, analysis, PPT/PDF/Excel export."""
import logging
import json
import copy
from datetime import date
from typing import Optional, Dict, List, Any, Literal
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import JSONResponse, StreamingResponse
from pydantic import BaseModel, Field, validator, root_validator
import pandas as pd
import io

from smartbi.services.financial_dashboard import FinancialDashboardService
from smartbi.services.financial.base import _sanitize_for_json
from smartbi.services.financial.export_conclusions import conclusion_rows as _conclusion_rows
from smartbi.services.factory_charts.export_guard import ExportGuardRoute, validate_export
from smartbi.services.factory_charts.details import canonical_period_label
from smartbi.services.factory_charts import conversation
from smartbi.services.factory_charts.cost_authorization import is_manufacturing_cost_column, row_product_ids

logger = logging.getLogger(__name__)
router = APIRouter(route_class=ExportGuardRoute)

# Service singleton
dashboard_service = FinancialDashboardService()


class FactoryRecommendationRequest(BaseModel):
    source: Literal['upload', 'business'] = 'business'
    upload_id: Optional[int] = Field(None, ge=1)
    inspection_upload_id: Optional[int] = Field(None, ge=1, alias='inspectionUploadId')
    cost_upload_id: Optional[int] = Field(None, ge=1, alias='costUploadId')
    factory_id: Optional[str] = None
    year: int = Field(2026, ge=1900, le=2100)
    period_type: Literal['all', 'year', 'month_range'] = 'all'
    start_month: int = Field(1, ge=1, le=12)
    end_month: int = Field(12, ge=1, le=12)
    include_optional: bool = True
    use_ai: bool = True
    guided: bool = False
    analysis_subject: Optional[str] = Field(None, max_length=80)
    # None keeps the legacy direct-route behaviour.  An empty list is used by
    # the AI catalogue pass, while a non-empty list lazily allocates only the
    # explicitly selected, already-authorized candidates.
    detail_candidate_ids: Optional[List[str]] = Field(None, max_items=4)

    class Config:
        allow_population_by_field_name = True
        populate_by_name = True

    @root_validator(pre=True)
    def reject_invalid_joint_ids(cls, values):
        source = values.get('source', 'business')
        keys = ('upload_id', 'inspection_upload_id', 'cost_upload_id',
                'uploadId', 'inspectionUploadId', 'costUploadId')
        for key in keys:
            if key in values and isinstance(values[key], bool):
                raise ValueError(f'{key}必须是正整数')
        related = values.get('inspection_upload_id', values.get('inspectionUploadId'))
        cost = values.get('cost_upload_id', values.get('costUploadId'))
        if source != 'upload' and (related is not None or cost is not None):
            raise ValueError('只有upload来源允许联合上传ID')
        return values

    @validator('detail_candidate_ids')
    def valid_detail_candidates(cls, value):
        if value is None:
            return value
        if any(not item or len(item) > 300 for item in value) or len(value) != len(set(value)):
            raise ValueError('详情图表标识必须非空、唯一且不超过300字符')
        return value

    @root_validator(skip_on_failure=True)
    def validate_uploaded_joint_ids(cls, values):
        inspection, cost, anchor = values.get('inspection_upload_id'), values.get('cost_upload_id'), values.get('upload_id')
        if (inspection is None) != (cost is None):
            raise ValueError('inspectionUploadId 与 costUploadId 必须同时提供')
        if inspection is not None and (anchor is None or len({anchor, inspection, cost}) != 3):
            raise ValueError('联合上传必须提供互不相同的 uploadId、inspectionUploadId、costUploadId')
        return values


class FactorySelectionRequest(BaseModel):
    factory_id: Optional[str] = None
    source: Literal['upload', 'business']
    selected_ids: List[str] = Field(..., max_items=100)

    @validator('selected_ids')
    def valid_selection(cls, value):
        if any(not item or len(item) > 300 for item in value) or len(value) != len(set(value)):
            raise ValueError('图表标识必须非空、唯一且不超过300字符')
        return value


class FactoryAnalysisRequest(BaseModel):
    """Free-question factory analysis request (public camelCase aliases)."""
    question: str = Field(..., min_length=1, max_length=1000)
    source: Optional[Literal['upload', 'business']] = None
    upload_id: Optional[int] = Field(None, ge=1, alias='uploadId')
    inspection_upload_id: Optional[int] = Field(None, ge=1, alias='inspectionUploadId')
    cost_upload_id: Optional[int] = Field(None, ge=1, alias='costUploadId')
    factory_id: Optional[str] = Field(None, alias='factoryId')
    year: Optional[int] = Field(None, ge=1900, le=2100)
    period_type: Optional[Literal['all', 'year', 'month_range']] = Field(None, alias='periodType')
    start_month: Optional[int] = Field(None, ge=1, le=12, alias='startMonth')
    end_month: Optional[int] = Field(None, ge=1, le=12, alias='endMonth')
    analysis_subject: Optional[str] = Field(None, max_length=300, alias='analysisSubject')
    context_id: Optional[str] = Field(None, max_length=200, alias='contextId')

    class Config:
        allow_population_by_field_name = True
        populate_by_name = True

    @root_validator(pre=True)
    def reject_invalid_joint_ids(cls, values):
        source = values.get('source', 'business')
        keys = ('upload_id', 'inspection_upload_id', 'cost_upload_id',
                'uploadId', 'inspectionUploadId', 'costUploadId')
        for key in keys:
            if key in values and isinstance(values[key], bool):
                raise ValueError(f'{key}必须是正整数')
        related = values.get('inspection_upload_id', values.get('inspectionUploadId'))
        cost = values.get('cost_upload_id', values.get('costUploadId'))
        if source != 'upload' and (related is not None or cost is not None):
            raise ValueError('只有upload来源允许联合上传ID')
        return values

    @root_validator(skip_on_failure=True)
    def validate_uploaded_joint_ids(cls, values):
        inspection, cost, anchor = values.get('inspection_upload_id'), values.get('cost_upload_id'), values.get('upload_id')
        if (inspection is None) != (cost is None):
            raise ValueError('inspectionUploadId 与 costUploadId 必须同时提供')
        if inspection is not None and (anchor is None or len({anchor, inspection, cost}) != 3):
            raise ValueError('联合上传必须提供互不相同的 uploadId、inspectionUploadId、costUploadId')
        return values

    @validator('question')
    def question_is_not_blank(cls, value):
        if not value.strip():
            raise ValueError('问题不能为空')
        return value.strip()


@router.get('/factory-selection')
async def load_factory_selection(http_request: Request, source: Literal['upload', 'business'], factory_id: Optional[str] = None):
    from smartbi.services.factory_charts.selection import access
    return await access(http_request, factory_id, source)


@router.put('/factory-selection')
async def save_factory_selection(request: FactorySelectionRequest, http_request: Request):
    from smartbi.services.factory_charts.selection import access
    return await access(http_request, request.factory_id, request.source, request.selected_ids)


@router.post('/factory-recommendations')
async def factory_recommendations(request: FactoryRecommendationRequest, http_request: Request):
    from common.tenant_identity import require_same_factory
    from smartbi_compat._rbac_role import ANALYTICS_READ_ROLES
    from smartbi.api.factory_production_gold import _can_view_cost
    from starlette.concurrency import run_in_threadpool
    from smartbi.services.factory_charts import engine, sources, ranking, finance_sources, operational_sources, captured_sources, trusted_business_sources, detail_store, details

    factory_id = require_same_factory(http_request, request.factory_id)
    role = getattr(http_request.state, 'role', None)
    if role not in ANALYTICS_READ_ROLES:
        raise HTTPException(403, '当前身份没有经营分析查看权限')
    if request.start_month > request.end_month:
        raise HTTPException(422, '开始月份不能晚于结束月份')
    if request.period_type == 'year' and (request.start_month != 1 or request.end_month != 12):
        raise HTTPException(422, '全年口径必须为1至12月')
    can_price = _can_view_cost(role)
    errors, extra_gaps = [], []
    if request.source == 'upload':
        if not request.upload_id:
            raise HTTPException(422, '请先选择本厂上传报表')
        if request.inspection_upload_id is not None and not can_price:
            raise HTTPException(403, '当前身份没有成本查看权限，不能读取联合成本上传')
        try:
            load_request = FinancialDashboardRequest(upload_id=request.upload_id)
            if request.inspection_upload_id is None:
                frame = await _load_data(load_request, factory_id, http_request)
            else:
                frame = await _load_data(load_request, factory_id, http_request,
                                         require_cost_authorization=True)
            if request.inspection_upload_id is not None:
                from smartbi.services.factory_charts.uploaded_joint import parse_selection, prepare_uploaded_join, dataset_for_join
                selection = parse_selection({'uploadId': request.upload_id,
                                             'inspectionUploadId': request.inspection_upload_id,
                                             'costUploadId': request.cost_upload_id})
                inspection = await _load_data(FinancialDashboardRequest(upload_id=selection.inspection_upload_id), factory_id,
                                               http_request, require_cost_authorization=True)
                cost = await _load_data(FinancialDashboardRequest(upload_id=selection.cost_upload_id), factory_id, http_request,
                                        require_cost_authorization=True)
                if len(frame) > 20000 or len(inspection) > 20000 or len(cost) > 20000:
                    raise HTTPException(422, '联合上传任一报表超过两万行安全上限，请拆分报表后重试')
                joined, join_gaps = prepare_uploaded_join(
                    engine.Dataset(frame, f'upload-{selection.production_upload_id}'),
                    engine.Dataset(inspection, f'upload-{selection.inspection_upload_id}'),
                    engine.Dataset(cost, f'upload-{selection.cost_upload_id}'))
                if joined is None:
                    return JSONResponse(status_code=422, content={'success': False, 'message': '上传数据无法形成完整生产-质检-成本联合分析', 'data': {'gaps': join_gaps, 'jointUploads': {'inspectionUploadId': selection.inspection_upload_id, 'costUploadId': selection.cost_upload_id}}})
                datasets = [dataset_for_join(selection, joined)]
                extra_gaps.extend(join_gaps)
            else:
                datasets = [engine.Dataset(frame, f'upload-{request.upload_id} · 本厂上传报表',
                                           engine.is_financial(frame), load_request.sheet_name or '')]
        except DataSourceError as e:
            return e.response()
        if len(frame) > 20000:
            raise HTTPException(422, '报表超过两万行安全上限，请拆分报表后重试')
    else:
        datasets, errors, extra_gaps = await sources.load(factory_id, request, can_price, auth_request=http_request)
        from smartbi.services.factory_charts.finance_auth import full_factory_finance_access
        finance_allowed = can_price and await full_factory_finance_access(http_request, factory_id)
        extra_data, extra_errors, more_gaps = await finance_sources.load(factory_id, request, finance_allowed)
        datasets.extend(extra_data)
        errors.extend(extra_errors)
        extra_gaps.extend(more_gaps)
        for loader in (operational_sources.load, captured_sources.load, trusted_business_sources.load):
            more_data, more_errors, source_gaps = await loader(factory_id, request, can_price)
            datasets.extend(more_data)
            errors.extend(more_errors)
            extra_gaps.extend(source_gaps)
        if not datasets and errors:
            return JSONResponse(status_code=503, content={'success': False, 'message': '工厂数据来源读取失败，请重试', 'data': {'sourceErrors': errors}})
    # The catalogue is deliberately built from the already-authorized datasets;
    # it never asks a broad source for product names.  Apply a selected subject
    # before chart generation so a missing subject cannot silently become a
    # full-factory answer.
    subjects = engine.analysis_subjects(datasets, can_price)
    scoped_datasets, subject_gap, scope_label = engine.scope_datasets(
        datasets, request, subjects, can_price)
    candidates, gaps = await run_in_threadpool(engine.build, scoped_datasets, request, can_price)
    if subject_gap:
        gaps.append(subject_gap)
    errors.extend({'source': gap['title'], 'message': gap['missing'][0]} for gap in gaps if gap.get('error'))
    gaps = [gap for gap in gaps if not gap.get('error')]
    candidates, order = await ranking.rank(candidates, request.use_ai, cache_scope=f'{factory_id}:{role}')
    # These three legacy gaps describe missing capture, not missing data in
    # unrelated datasets. Remove each only after its trusted business adapter
    # has produced the corresponding real candidate for this period.
    proven = set()
    for candidate in candidates:
        source = str(candidate.get('source', '')).split(' · ', 1)[0]
        metric = str(candidate.get('id', ''))
        if source == 'sales-sales_channel' and 'finance.channel_analysis' in metric:
            proven.add('sales.channel')
        if source in ('sales-department_id', 'department') and 'finance.business_unit_analysis' in metric:
            proven.add('finance.unit-source')
        if source in ('oee', 'oee-sheet') and 'production.oee' in metric:
            proven.add('production.oee-source')
    extra_gaps = [gap for gap in extra_gaps if gap.get('id') not in proven]
    from common.tenant_identity import caller_user_id
    user_id = caller_user_id(http_request)
    if request.detail_candidate_ids is not None:
        authorized_by_id = {str(item.get('id')): item for item in candidates if item.get('id')}
        requested_ids = list(request.detail_candidate_ids)
        unknown_ids = [item_id for item_id in requested_ids if item_id not in authorized_by_id]
        if unknown_ids:
            raise HTTPException(422, '详情图表不在当前授权目录中')
        detail_candidates = [authorized_by_id[item_id] for item_id in requested_ids]
    else:
        detail_candidates = candidates
    if user_id is not None and user_id > 0 and request.detail_candidate_ids == []:
        # Discovery remains stateless: prepare the same opaque immutable
        # snapshot proof, but do not consume a shared context or touch PG.
        candidates = detail_store.prepare_shared_metadata(
            candidates, scoped_datasets, request, (factory_id, user_id, role), can_price)
    elif user_id is not None and user_id > 0 and (request.detail_candidate_ids is None or detail_candidates):
        try:
            registered = await detail_store.register_shared(detail_candidates, scoped_datasets,
                                                             request, (factory_id, user_id, role), can_price)
            if request.detail_candidate_ids is None:
                candidates = registered
            else:
                registered_by_id = {str(item.get('id')): item for item in registered if item.get('id')}
                candidates = [registered_by_id.get(str(item.get('id')), item) for item in candidates]
        except details.DetailError as exc:
            raise HTTPException(exc.status_code, exc.detail) from None
    from smartbi.services.factory_charts.selection_groups import default_selection
    label = canonical_period_label(request)
    return _sanitize_for_json({'success': True, 'data': {
        'factoryId': factory_id, 'source': request.source,
        'uploadId': request.upload_id if request.source == 'upload' else None,
        'jointUploads': ({'inspectionUploadId': request.inspection_upload_id,
                          'costUploadId': request.cost_upload_id}
                         if request.inspection_upload_id is not None else None),
        'sourceLabel': '本厂业务系统' if request.source == 'business' else f'本厂上传报表 #{request.upload_id}',
        'period': {'label': label, 'year': request.year, 'periodType': request.period_type,
                   'startMonth': request.start_month, 'endMonth': request.end_month}, 'ranking': order,
        'recommendations': candidates, 'defaultSelection': default_selection(candidates, guided=request.guided),
        'subjects': subjects,
        'analysisSubject': request.analysis_subject,
        'analysisScopeLabel': scope_label,
        'unavailable': gaps + extra_gaps,
        'sourceErrors': errors,
    }})


def _recommendation_body(response: Any) -> dict[str, Any]:
    """Read the existing recommendation route's JSON without changing its contract."""
    if isinstance(response, JSONResponse):
        status = int(response.status_code)
        raw = response.body
        payload = json.loads(raw.decode('utf-8') if isinstance(raw, bytes) else raw)
        if status >= 400:
            if isinstance(payload.get('data'), dict) and payload['data'].get('sourceErrors'):
                return payload
            raise HTTPException(status, payload.get('message', '工厂图表数据暂不可用'))
        return payload
    if isinstance(response, dict):
        return response
    raise HTTPException(503, '工厂图表数据暂不可用')


def _recommendation_request(source: str, upload_id: Optional[int], factory_id: str,
                            scope: dict[str, Any], *, guided: bool = False,
                            analysis_subject: Optional[str] = None,
                            detail_candidate_ids: Optional[List[str]] = (),
                            inspection_upload_id: Optional[int] = None,
                            cost_upload_id: Optional[int] = None) -> FactoryRecommendationRequest:
    selection = scope.get('uploadedJoinSelection') if isinstance(scope, dict) else None
    if isinstance(selection, dict):
        inspection_upload_id = inspection_upload_id or selection.get('inspectionUploadId')
        cost_upload_id = cost_upload_id or selection.get('costUploadId')
    return FactoryRecommendationRequest(
        source=source, upload_id=upload_id, factory_id=factory_id,
        year=int(scope.get('year', 2026)),
        period_type=scope.get('period_type', 'all'),
        start_month=int(scope.get('start_month', 1)),
        end_month=int(scope.get('end_month', 12)),
        include_optional=True, use_ai=False, guided=guided,
        analysis_subject=analysis_subject,
        detail_candidate_ids=detail_candidate_ids,
        inspection_upload_id=inspection_upload_id,
        cost_upload_id=cost_upload_id,
    )


def _analysis_response(data: dict[str, Any], *, factory_id: str, source: str,
                      question: str, status: str, answer: str, selected_ids: list[str],
                      recommendations: list[dict[str, Any]], catalog: list[dict[str, Any]],
                      facts: list[dict[str, Any]], limitations: list[str],
                      context_id: str, subjects: list[dict[str, Any]],
                      follow_up_questions: list[str]) -> dict[str, Any]:
    # Keep the public response stable and camelCase while selected
    # recommendations remain byte-for-byte the trusted legacy candidates.
    return _sanitize_for_json({'success': True, 'data': {
        'factoryId': factory_id,
        'source': source,
        'uploadId': data.get('uploadId'),
        'sourceLabel': data.get('sourceLabel'),
        'jointUploads': data.get('jointUploads'),
        'period': data.get('period', {}),
        'analysisSubject': data.get('analysisSubject'),
        'analysisScopeLabel': data.get('analysisScopeLabel'),
        'question': question,
        'status': status,
        'answer': answer,
        'selectedIds': selected_ids,
        'recommendations': recommendations,
        'catalog': catalog,
        'facts': facts,
        'limitations': limitations,
        'sourceErrors': data.get('sourceErrors') or [],
        'followUpQuestions': follow_up_questions,
        'contextId': context_id,
        'subjects': subjects,
    }, 'message': answer if status != 'answered' else '工厂分析完成'})


@router.post('/factory-analysis')
async def factory_analysis(request: FactoryAnalysisRequest, http_request: Request):
    """Grounded free-question chart planning for the authenticated factory."""
    from common.tenant_identity import require_same_factory

    factory_id = require_same_factory(http_request, request.factory_id)
    owner = conversation._owner(http_request)
    try:
        context = await conversation.claim_shared_context(request.context_id, owner) if request.context_id else None
    except Exception as exc:
        if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
            raise HTTPException(exc.status_code, exc.detail) from None
        raise
    if context is not None:
        # Serialize a turn from context lookup through planner validation.  A
        # second request cannot mix its source/scope with an in-flight turn.
        try:
            async with context.lock:
                response = await _factory_analysis_turn(request, http_request, factory_id, owner, context)
            try:
                await conversation.save_shared_context(context)
            except Exception as exc:
                if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
                    raise HTTPException(exc.status_code, exc.detail) from None
                raise
            return response
        except Exception:
            await conversation.release_shared_context(context)
            raise

    source = request.source or 'business'
    upload_id = request.upload_id
    current_year = date.today().year
    scope = {
        'year': request.year or current_year,
        'period_type': request.period_type or 'all',
        'start_month': request.start_month or 1,
        'end_month': request.end_month or 12,
        'guided': True,
        'analysis_subject': None,
    }
    if request.inspection_upload_id is not None:
        scope['uploadedJoinSelection'] = {
            'uploadId': upload_id,
            'inspectionUploadId': request.inspection_upload_id,
            'costUploadId': request.cost_upload_id,
        }
    explicit_scope, scope_error, explicit_present = conversation.apply_explicit_period(
        scope, year=request.year, period_type=request.period_type,
        start_month=request.start_month, end_month=request.end_month)
    if scope_error:
        raise HTTPException(422, scope_error)
    scope = explicit_scope
    rec_request = _recommendation_request(source, upload_id, factory_id, scope, guided=True,
                                           inspection_upload_id=request.inspection_upload_id,
                                           cost_upload_id=request.cost_upload_id)
    recommendation_data = _recommendation_body(await factory_recommendations(rec_request, http_request))['data']
    subjects = list(recommendation_data.get('subjects') or [])
    resolved_subject, subject_error = conversation.resolve_subject(request.question, subjects, request.analysis_subject)
    # Natural-language scope applies on the first turn as well as follow-ups.
    # Resolve it against the fresh authorized catalogue before returning or
    # invoking the planner; otherwise "去年/上个月" silently stayed current/all.
    period_question = request.question
    if resolved_subject:
        for item in subjects:
            if str(item.get('id')) == resolved_subject:
                period_question = period_question.replace(str(item.get('label', '')), ' ')
                break
    natural_scope, natural_error = conversation.resolve_period(period_question, scope)
    hard_explicit = (request.period_type in {'year', 'month_range'} or
                     (request.period_type is None and any(value is not None for value in (request.year, request.start_month, request.end_month))))
    if not natural_error and hard_explicit and natural_scope != scope:
        natural_error = '问题中的自然语言期间与显式期间不一致，请只指定一个期间'
    if natural_error:
        try:
            context = await conversation.create_shared_context(owner, source, upload_id, scope)
        except Exception as exc:
            if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
                raise HTTPException(exc.status_code, exc.detail) from None
            raise
        try:
            response = await _render_factory_analysis_turn(request, http_request, factory_id, context,
                recommendation_data, subjects, None, period_error=natural_error, scope=scope)
            try:
                await conversation.save_shared_context(context)
            except Exception as exc:
                if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
                    raise HTTPException(exc.status_code, exc.detail) from None
                raise
            return response
        except Exception:
            await conversation.release_shared_context(context)
            raise
    if natural_scope != scope:
        scope = natural_scope
        rec_request = _recommendation_request(source, upload_id, factory_id, scope, guided=True,
                                               analysis_subject=resolved_subject)
        recommendation_data = _recommendation_body(await factory_recommendations(rec_request, http_request))['data']
        subjects = list(recommendation_data.get('subjects') or subjects)
    unresolved = bool(subject_error)
    if resolved_subject:
        scope['analysis_subject'] = resolved_subject
        rec_request = _recommendation_request(source, upload_id, factory_id, scope, guided=True,
                                               analysis_subject=resolved_subject)
        recommendation_data = _recommendation_body(await factory_recommendations(rec_request, http_request))['data']
        subjects = list(recommendation_data.get('subjects') or subjects)
    try:
        context = await conversation.create_shared_context(owner, source, upload_id, scope,
                                                           question=request.question, unresolved_product=unresolved)
    except Exception as exc:
        if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
            raise HTTPException(exc.status_code, exc.detail) from None
        raise
    try:
        response = await _render_factory_analysis_turn(request, http_request, factory_id, context,
            recommendation_data, subjects, subject_error)
        try:
            await conversation.save_shared_context(context)
        except Exception as exc:
            if hasattr(exc, 'status_code') and hasattr(exc, 'detail'):
                raise HTTPException(exc.status_code, exc.detail) from None
            raise
        return response
    except Exception:
        await conversation.release_shared_context(context)
        raise


async def _factory_analysis_turn(request: FactoryAnalysisRequest, http_request: Request,
                                 factory_id: str, owner: tuple[str, int, str],
                                 context: conversation.ConversationContext):
    if request.source is not None and request.source != context.source:
        raise HTTPException(403, '对话上下文数据源不可切换')
    if request.upload_id is not None and request.upload_id != context.upload_id:
        raise HTTPException(403, '对话上下文上传数据源不可切换')
    # Every turn re-enters the old authorized recommendation route, so current
    # price permission, source errors, and changed business data are rechecked.
    base_scope = dict(context.scope)
    base_scope['guided'] = True
    base_request = _recommendation_request(context.source, context.upload_id, factory_id, base_scope,
                                           guided=True, analysis_subject=base_scope.get('analysis_subject'))
    recommendation_data = _recommendation_body(await factory_recommendations(base_request, http_request))['data']
    subjects = list(recommendation_data.get('subjects') or [])
    # Product resolution must use this turn's authorized catalogue; retaining
    # only subject metadata in the context is not permission to reuse it.
    current_subject_ids = {str(item.get('id')) for item in subjects if item.get('id')}
    if context.scope.get('analysis_subject') and context.scope.get('analysis_subject') not in current_subject_ids:
        context.scope = {**context.scope, 'analysis_subject': None, 'guided': True}
        context.unresolved_product = True
    scope, _, period_error = conversation.scope_for_context(
        context, request.question, request.analysis_subject, subjects,
        year=request.year, period_type=request.period_type,
        start_month=request.start_month, end_month=request.end_month)
    # Resolve against the current authorized product catalogue, never a stale
    # context copy.  An unresolved turn cannot fall back to factory scope.
    _, subject_error = conversation.resolve_subject(request.question, subjects, request.analysis_subject)
    if period_error:
        conversation.touch_context(context, request.question, [], unresolved_product=False,
                                    unresolved_period=True, period_error=period_error)
        return await _render_factory_analysis_turn(request, http_request, factory_id, context,
                                                    recommendation_data, [], None, period_error=period_error, scope=scope)
    if subject_error:
        scope_error = subject_error
        if conversation.is_cancel(request.question):
            conversation.touch_context(context, request.question, [], unresolved_product=False)
            return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                      question=request.question, status='clarification',
                                      answer='已取消本轮产品限定分析。若要继续，请明确指定产品。', selected_ids=[],
                                      recommendations=[], catalog=conversation.semantic_catalog(recommendation_data.get('recommendations') or []),
                                      facts=[], limitations=[scope_error], context_id=context.context_id,
                                      subjects=subjects, follow_up_questions=['请指定一个明确的产品名称或SKU'])
        conversation.touch_context(context, request.question, [], unresolved_product=True,
                                   unresolved_period=False, period_error=None)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='clarification',
                                  answer=scope_error + '。请明确一个当前工厂产品，或明确取消。', selected_ids=[],
                                  recommendations=[], catalog=conversation.semantic_catalog(recommendation_data.get('recommendations') or []),
                                  facts=[], limitations=[scope_error], context_id=context.context_id,
                                  subjects=subjects, follow_up_questions=['请指定一个明确的产品名称或SKU', '取消本轮分析'])
    # An unsupported period must be corrected before a bare follow-up can
    # broaden to the saved baseline.  A valid explicit/NL period clears it.
    if context.unresolved_period and scope == context.scope:
        conversation.touch_context(context, request.question, [], unresolved_product=False,
                                   unresolved_period=True, period_error=context.period_error)
        return await _render_factory_analysis_turn(request, http_request, factory_id, context,
                                                    recommendation_data, [], None,
                                                    period_error=context.period_error, scope=scope)
    if context.unresolved_product and not scope.get('analysis_subject'):
        if conversation.is_cancel(request.question):
            conversation.touch_context(context, request.question, [], unresolved_product=False)
            return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                      question=request.question, status='clarification',
                                      answer='已取消本轮产品限定分析。', selected_ids=[], recommendations=[],
                                      catalog=conversation.semantic_catalog(recommendation_data.get('recommendations') or []),
                                      facts=[], limitations=['本轮分析已取消'], context_id=context.context_id,
                                      subjects=subjects, follow_up_questions=['如需继续，请指定一个明确的产品名称或SKU'])
        conversation.touch_context(context, request.question, [], unresolved_product=True)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='clarification',
                                  answer='上一轮产品尚未明确，不能继续扩大到全厂。请指定产品或明确取消。', selected_ids=[],
                                  recommendations=[], catalog=conversation.semantic_catalog(recommendation_data.get('recommendations') or []),
                                  facts=[], limitations=['产品范围未解析'], context_id=context.context_id,
                                  subjects=subjects, follow_up_questions=['请指定一个明确的产品名称或SKU', '取消本轮分析'])
    if scope.get('analysis_subject') and scope.get('analysis_subject') != context.scope.get('analysis_subject'):
        final_request = _recommendation_request(context.source, context.upload_id, factory_id, scope,
                                                guided=True, analysis_subject=scope.get('analysis_subject'))
        recommendation_data = _recommendation_body(await factory_recommendations(final_request, http_request))['data']
        subjects = list(recommendation_data.get('subjects') or subjects)
    elif scope != context.scope:
        final_request = _recommendation_request(context.source, context.upload_id, factory_id, scope, guided=True,
                                                analysis_subject=scope.get('analysis_subject'))
        recommendation_data = _recommendation_body(await factory_recommendations(final_request, http_request))['data']
        subjects = list(recommendation_data.get('subjects') or subjects)
    return await _render_factory_analysis_turn(request, http_request, factory_id, context,
                                                recommendation_data, subjects, None,
                                                period_error=None, scope=scope)


async def _render_factory_analysis_turn(request: FactoryAnalysisRequest, http_request: Request,
                                        factory_id: str, context: conversation.ConversationContext,
                                        recommendation_data: dict[str, Any], subjects: list[dict[str, Any]],
                                        subject_error: Optional[str],
                                        period_error: Optional[str] = None,
                                        scope: Optional[dict[str, Any]] = None):
    catalog = conversation.semantic_catalog(recommendation_data.get('recommendations') or [])
    effective_scope = scope or context.scope
    if conversation.is_joint_intent(request.question) and not effective_scope.get('analysis_subject'):
        conversation.touch_context(context, request.question, [], unresolved_product=True,
                                    unresolved_period=False, period_error=None, scope=effective_scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='clarification',
                                  answer='联合/多维度分析需要先明确一个当前工厂产品，不能把独立图表冒充联合结果。',
                                  selected_ids=[], recommendations=[], catalog=catalog, facts=[],
                                  limitations=['请先指定产品后再验证生产、质量、良品率与成本的真实联合关系'],
                                  context_id=context.context_id, subjects=subjects,
                                  follow_up_questions=['请指定一个明确的产品名称或SKU'])
    if period_error:
        conversation.touch_context(context, request.question, [], unresolved_product=False,
                                    unresolved_period=True, period_error=period_error, scope=scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='clarification',
                                  answer=period_error + '。请明确年份或月份区间。', selected_ids=[],
                                  recommendations=[], catalog=catalog, facts=[], limitations=[period_error],
                                  context_id=context.context_id, subjects=[],
                                  follow_up_questions=['请明确年份或月份区间', '例如：2025年或2026年8月至10月'])
    if subject_error:
        conversation.touch_context(context, request.question, [], unresolved_product=True,
                                   unresolved_period=False, period_error=None, scope=scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='clarification',
                                  answer=subject_error + '。请明确一个当前工厂产品，或明确取消。', selected_ids=[],
                                  recommendations=[], catalog=catalog, facts=[], limitations=[subject_error],
                                  context_id=context.context_id, subjects=subjects,
                                  follow_up_questions=['请指定一个明确的产品名称或SKU', '取消本轮分析'])
    selected_ids, planner_mode = await conversation.plan(
        request.question, catalog, context.history, factory_id=factory_id,
        authorized_subjects=subjects,
        authorized_recommendations=recommendation_data.get('recommendations') or [])
    if planner_mode.get('mode') != 'ai':
        conversation.touch_context(context, request.question, [], scope=scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='unavailable',
                                  answer='AI图表规划暂不可用，未返回未经验证的图表或事实。', selected_ids=[],
                                  recommendations=[], catalog=catalog, facts=[],
                                  limitations=[planner_mode.get('reason', 'AI规划不可用')],
                                  context_id=context.context_id, subjects=subjects,
                                  follow_up_questions=['请稍后重试，或先从图表目录选择分析主题'])
    by_id = conversation._catalog_by_id(recommendation_data.get('recommendations') or [])
    if any(item_id not in by_id for item_id in selected_ids):
        conversation.touch_context(context, request.question, [], scope=scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='unavailable',
                                  answer='当前授权数据已变化，AI选择的图表无法重新验证。', selected_ids=[],
                                  recommendations=[], catalog=catalog, facts=[],
                                  limitations=['AI选择不在当前授权目录中'], context_id=context.context_id,
                                  subjects=subjects, follow_up_questions=['请重新提问'])
    # The planner only saw a metadata catalogue.  Re-read the same authorized
    # source/scope and lazily allocate contexts for its actual 1-4 selections.
    # This second pass also prevents stale IDs, source changes, or unit changes
    # from gaining a detail context through an old recommendation response.
    selected_request = _recommendation_request(
        context.source, context.upload_id, factory_id, effective_scope,
        guided=True, analysis_subject=effective_scope.get('analysis_subject'),
        detail_candidate_ids=selected_ids,
    )
    selected_data = _recommendation_body(
        await factory_recommendations(selected_request, http_request)
    )['data']
    if selected_data.get('period') != recommendation_data.get('period') \
            or selected_data.get('source') != recommendation_data.get('source'):
        conversation.touch_context(context, request.question, [], scope=scope)
        return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                  question=request.question, status='unavailable',
                                  answer='当前授权数据已变化，AI选择的图表无法重新验证。', selected_ids=[],
                                  recommendations=[], catalog=catalog, facts=[],
                                  limitations=['AI选择的图表来源或范围已变化'],
                                  context_id=context.context_id, subjects=subjects,
                                  follow_up_questions=['请重新提问'])
    fresh_by_id = conversation._catalog_by_id(selected_data.get('recommendations') or [])
    for item_id in selected_ids:
        old_item = by_id[item_id]
        fresh_item = fresh_by_id.get(item_id)
        if fresh_item is None or str(fresh_item.get('source', '')) != str(old_item.get('source', '')) \
                or str(fresh_item.get('unit', '')) != str(old_item.get('unit', '')):
            conversation.touch_context(context, request.question, [], scope=scope)
            return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                      question=request.question, status='unavailable',
                                      answer='当前授权数据已变化，AI选择的图表无法重新验证。', selected_ids=[],
                                      recommendations=[], catalog=catalog, facts=[],
                                      limitations=['AI选择的图表来源、范围或单位已变化'],
                                      context_id=context.context_id, subjects=subjects,
                                      follow_up_questions=['请重新提问'])
    recommendation_data = selected_data
    catalog = conversation.semantic_catalog(recommendation_data.get('recommendations') or [])
    by_id = fresh_by_id
    selected = [by_id[item_id] for item_id in selected_ids]
    if conversation.is_joint_intent(request.question):
        required_family = conversation.joint_family_requirement(request.question)
        fused_ids = {str(item.get('id')) for item in catalog
                     if str(item.get('familyId', '')) == required_family}
        if not fused_ids:
            conversation.touch_context(context, request.question, [], unresolved_product=False,
                                       unresolved_period=False, period_error=None, scope=scope)
            return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                      question=request.question, status='unavailable',
                                      answer='当前授权数据没有可验证的生产-质量-成本联合图表，未将独立图表冒充联合分析。',
                                      selected_ids=[], recommendations=[], catalog=catalog, facts=[],
                                      limitations=[f'缺少满足同厂、同产品、同单位且真实对齐的 {required_family} 联合关系'],
                                      context_id=context.context_id, subjects=subjects,
                                      follow_up_questions=['请稍后重试或改为单一维度分析'])
        if not any(item_id in fused_ids for item_id in selected_ids):
            conversation.touch_context(context, request.question, [], unresolved_product=False,
                                       unresolved_period=False, period_error=None, scope=scope)
            return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                                      question=request.question, status='unavailable',
                                      answer='当前规划未选择已验证的联合图表，未将独立图表冒充联合分析。',
                                      selected_ids=[], recommendations=[], catalog=catalog, facts=[],
                                      limitations=[f'联合分析必须由 {required_family} 真实候选支撑'],
                                      context_id=context.context_id, subjects=subjects,
                                      follow_up_questions=['请重新提问联合生产、质量与成本'])
    facts = conversation.facts_for(selected)
    limitations = conversation.limitations_for(selected, catalog, request.question)
    if conversation.is_joint_intent(request.question):
        required_family = conversation.joint_family_requirement(request.question)
        for gap in recommendation_data.get('unavailable') or []:
            if isinstance(gap, dict) and str(gap.get('id')) == required_family:
                missing = gap.get('missing') if isinstance(gap.get('missing'), list) else []
                limitations.extend(str(item) for item in missing[:8] if item)
        limitations = list(dict.fromkeys(limitations))[:16]
    answer = conversation.render_answer(request.question, selected, facts, limitations)
    conversation.touch_context(context, request.question, selected_ids, unresolved_product=False,
                               unresolved_period=False, period_error=None, scope=scope)
    return _analysis_response(recommendation_data, factory_id=factory_id, source=context.source,
                              question=request.question, status='answered', answer=answer,
                              selected_ids=selected_ids, recommendations=selected, catalog=catalog,
                              facts=facts, limitations=limitations, context_id=context.context_id,
                              subjects=subjects,
                              follow_up_questions=['再看质量指标', '换一个分析维度', '查看更多图表'])


class FinancialDashboardRequest(BaseModel):
    upload_id: Optional[int] = None
    raw_data: Optional[List[Dict[str, Any]]] = Field(None, max_items=50000)  # Direct data input, max 50k rows
    chart_type: str = "all"  # specific type, "all", or "auto"
    year: int = 2026
    period_type: str = "year"  # month|quarter|year|month_range
    start_month: int = Field(1, ge=1, le=12)
    end_month: int = Field(12, ge=1, le=12)
    # 租户以登录身份为准; 只有不带工厂的身份 (平台管理员) 才需要显式传。
    factory_id: Optional[str] = None
    filters: Optional[Dict[str, List[str]]] = None  # Fix 72: Global slicer filters {col: [values]}
    # raw_data 调用方可以传工作表名; upload_id 时以库里记录的工作表名为准
    sheet_name: Optional[str] = None


class AnalyzeRequest(BaseModel):
    chart_type: str
    analysis_context: str
    chart_result: Optional[Dict] = None
    follow_up_question: Optional[str] = None
    previous_analysis: Optional[str] = None


class PPTExportRequest(BaseModel):
    _export_limits = root_validator(pre=True, allow_reuse=True)(validate_export)
    chart_titles: Dict[str, str] = Field(default_factory=dict)
    period_label: Optional[str] = None
    report_title: Optional[str] = None
    upload_id: Optional[int] = None
    year: int = 2026
    period_type: str = "year"
    start_month: int = 1
    end_month: int = 12
    chart_images: Dict[str, str] = Field(default_factory=dict)  # {chartType: base64PNG}
    headline: Optional[Dict[str, Any]] = None
    findings: Optional[Dict[str, List[str]]] = None
    analysis_results: Dict[str, str] = Field(default_factory=dict)
    template: str = "default"
    company_name: str = "白垩纪科技"
    kpi_summary: Optional[Dict] = None


class DataSourceError(Exception):
    """数据源不可用: 带 HTTP 状态码与给用户看的原文。"""

    def __init__(self, status_code: int, message: str):
        super().__init__(message)
        self.status_code = status_code
        self.message = message

    def response(self) -> JSONResponse:
        return JSONResponse(status_code=self.status_code,
                            content={"success": False, "message": self.message, "error": self.message})


def _resolve_factory_id(http_request: Request, requested: Optional[str]) -> str:
    """租户取自登录身份 (JWTAuthMiddleware 写入 request.state); 请求体不能改写它。"""
    token_factory = getattr(http_request.state, "factory_id", None)
    if token_factory:
        if requested and requested != token_factory:
            logger.warning("financial-dashboard factory mismatch: token=%s body=%s", token_factory, requested)
            raise DataSourceError(403, "无权访问其他工厂的数据")
        return token_factory
    if requested:
        # 不带工厂的身份 (平台管理员 / 内部调用未带 X-Factory-Id) 必须显式指定
        return requested
    raise DataSourceError(400, "无法确定当前工厂，请重新登录")


# Resolve aliases from jsonb_object_keys so bracketed/parenthesized engine
# labels (for example 产品ID（单位）) use the same canonical identity as the
# in-memory engine.  Only keys whose normalized name is a stable ID are read;
# cost values are never enumerated as metadata.
_UPLOAD_ID_KEY_NORM = "lower(regexp_replace(k.key_name, '[（(][^（）()]*[）)]', '', 'g'))"
_UPLOAD_ID_NAMES_SQL = "'产品id','产品编号','产品编码','sku','sku编码','product_id','productid','product_type_id','producttypeid','产品类型id'"
_UPLOAD_ID_EXPR = "(SELECT NULLIF(BTRIM(d.row_data->>k.key_name),'') FROM jsonb_object_keys(d.row_data) AS k(key_name) WHERE {_norm} IN ({_names}) ORDER BY k.key_name LIMIT 1)".format(_norm=_UPLOAD_ID_KEY_NORM, _names=_UPLOAD_ID_NAMES_SQL)
_UPLOAD_ID_CONFLICT_EXPR = "(SELECT COUNT(DISTINCT BTRIM(d.row_data->>k.key_name)) FROM jsonb_object_keys(d.row_data) AS k(key_name) WHERE {_norm} IN ({_names}) AND BTRIM(d.row_data->>k.key_name) <> '') > 1".format(_norm=_UPLOAD_ID_KEY_NORM, _names=_UPLOAD_ID_NAMES_SQL)


async def _load_data(request: FinancialDashboardRequest, factory_id: Optional[str], auth_request=None,
                     require_cost_authorization: bool = False) -> pd.DataFrame:
    """Load data from raw_data, or from the uploaded rows of upload_id (one row_data per Excel row).

    upload_id 时把库里的工作表名写回 request.sheet_name (工作表名用于推断无科目表的口径)。
    """
    if request.raw_data:
        return pd.DataFrame(request.raw_data)

    if not request.upload_id:
        raise DataSourceError(400, "请先选择已上传的数据源")

    # 模块属性访问而不是 from-import: 测试与运行期替换连接池都要能生效
    import smartbi.config as smartbi_config
    try:
        pool = await smartbi_config.get_pg_pool()
        if pool is None:
            raise DataSourceError(503, "SmartBI 数据库未配置，无法读取上传数据")
        async with pool.acquire() as conn:
            params = [request.upload_id, factory_id]
            product_ids = []
            cost_keys = []
            hidden_cost_keys = []
            rows_preloaded = False
            low_cost_role = False
            authorized_product_ids = set()
            unscoped_authorized = False
            # Metadata-only pass for every stored upload.  It reads only stable
            # product IDs and JSON key names; row values are fetched below.
            id_rows = await conn.fetch(
                    """
                    SELECT DISTINCT {_id_expr} AS product_id,
                        {_conflict} AS id_conflict,
                        d.row_data->>'productId' AS "productId", d.row_data->>'产品ID' AS "产品ID",
                        d.row_data->>'产品编号' AS "产品编号", d.row_data->>'产品编码' AS "产品编码",
                        d.row_data->>'product_id' AS "product_id", d.row_data->>'productTypeId' AS "productTypeId",
                        d.row_data->>'product_type_id' AS "product_type_id", d.row_data->>'产品类型id' AS "产品类型id",
                        d.row_data->>'SKU' AS "SKU", d.row_data->>'sku' AS "sku", d.row_data->>'sku编码' AS "sku编码",
                        key_name
                    FROM smart_bi_dynamic_data d
                    JOIN smart_bi_pg_excel_uploads u ON u.id = d.upload_id
                    LEFT JOIN LATERAL jsonb_object_keys(d.row_data) AS keys(key_name) ON TRUE
                    WHERE d.upload_id = $1 AND u.factory_id = $2 AND d.factory_id = $2
                      AND u.upload_status = 'COMPLETED'
                    """.format(_id_expr=_UPLOAD_ID_EXPR, _conflict=_UPLOAD_ID_CONFLICT_EXPR), request.upload_id, factory_id)
            product_ids = sorted({value for row in id_rows for value in row_product_ids(dict(row))[0]})
            has_unidentified_cost_rows = any(not row_product_ids(dict(row))[0] or bool(row.get('id_conflict')) for row in id_rows
                                             if is_manufacturing_cost_column(row.get('key_name')))
            cost_keys = sorted({str(row['key_name']) for row in id_rows
                                if is_manufacturing_cost_column(row.get('key_name'))})
            if cost_keys:
                if auth_request is None:
                    raise DataSourceError(403, '无法验证上传制造成本授权')
                from smartbi.api.factory_production_gold import _can_view_cost
                role = getattr(getattr(auth_request, 'state', None), 'role', None)
                if not _can_view_cost(role):
                    low_cost_role = True
                    hidden_cost_keys = list(cost_keys)
                    rows = await conn.fetch(
                        """
                        SELECT CASE WHEN FALSE THEN d.row_data ELSE d.row_data - $3::text[] END AS row_data, d.sheet_name
                        FROM smart_bi_dynamic_data d
                        JOIN smart_bi_pg_excel_uploads u ON u.id = d.upload_id
                        WHERE d.upload_id = $1 AND u.factory_id = $2 AND d.factory_id = $2
                          AND u.upload_status = 'COMPLETED'
                        ORDER BY d.row_index
                        """, request.upload_id, factory_id, cost_keys)
                    cost_keys = []
                    rows_preloaded = True
            if cost_keys:
                from smartbi.services.factory_charts.cost_authorization import authorize_batches as authorize, CostAuthorizationError
                try:
                    authority = await authorize(auth_request, factory_id, product_ids,
                                                unscoped=(not product_ids or has_unidentified_cost_rows))
                except CostAuthorizationError:
                    raise DataSourceError(403, '当前授权不能读取上传制造成本')
                allowed = sorted(authority.allowed_product_ids)
                authorized_product_ids = set(allowed)
                unscoped_authorized = authority.unscoped_cost_allowed
                if not product_ids and authority.unscoped_cost_allowed:
                    rows = await conn.fetch(
                        """
                        SELECT d.row_data, d.sheet_name
                        FROM smart_bi_dynamic_data d
                        JOIN smart_bi_pg_excel_uploads u ON u.id = d.upload_id
                        WHERE d.upload_id = $1 AND u.factory_id = $2 AND d.factory_id = $2
                          AND u.upload_status = 'COMPLETED'
                        ORDER BY d.row_index
                        """, request.upload_id, factory_id)
                else:
                    rows = await conn.fetch(
                    """
                    SELECT CASE WHEN NOT {_conflict} AND ({_id_expr} = ANY($3::text[])
                                OR ({_id_expr} IS NULL AND $5::boolean))
                                THEN d.row_data ELSE d.row_data - $4::text[] END AS row_data, d.sheet_name
                    FROM smart_bi_dynamic_data d
                    JOIN smart_bi_pg_excel_uploads u ON u.id = d.upload_id
                    WHERE d.upload_id = $1 AND u.factory_id = $2 AND d.factory_id = $2
                      AND u.upload_status = 'COMPLETED'
                    ORDER BY d.row_index
                        """.format(_conflict=_UPLOAD_ID_CONFLICT_EXPR, _id_expr=_UPLOAD_ID_EXPR),
                        request.upload_id, factory_id, allowed, cost_keys, authority.unscoped_cost_allowed)
            else:
                if rows_preloaded:
                    pass
                else:
                    rows = await conn.fetch(
                    """
                    SELECT d.row_data, d.sheet_name
                    FROM smart_bi_dynamic_data d
                    JOIN smart_bi_pg_excel_uploads u ON u.id = d.upload_id
                    WHERE d.upload_id = $1 AND u.factory_id = $2 AND d.factory_id = $2
                      AND u.upload_status = 'COMPLETED'
                    ORDER BY d.row_index
                    """, request.upload_id, factory_id)
    except DataSourceError:
        raise
    except Exception as e:
        logger.error(f"Failed to load upload {request.upload_id} for {factory_id}: {e}", exc_info=True)
        raise DataSourceError(500, "读取上传数据失败，请稍后重试")

    if not rows:
        raise DataSourceError(404, f"未找到数据源 {request.upload_id}，或它不属于当前工厂")

    import json
    records = [json.loads(r["row_data"]) if isinstance(r["row_data"], str) else r["row_data"] for r in rows]
    sheet_names = {r["sheet_name"] for r in rows if r["sheet_name"]}
    if len(sheet_names) == 1:
        request.sheet_name = sheet_names.pop()
    frame = pd.DataFrame(records)
    if low_cost_role:
        frame = frame.drop(columns=[column for column in frame.columns
                                    if column in hidden_cost_keys or is_manufacturing_cost_column(column)], errors='ignore')
    elif cost_keys and authorized_product_ids is not None:
        for index, record in enumerate(records):
            ids, conflict = row_product_ids(dict(record))
            permitted = (not conflict and bool(ids) and ids <= authorized_product_ids) or (
                not conflict and not ids and unscoped_authorized)
            if not permitted:
                for column in cost_keys:
                    if column in frame.columns:
                        frame.loc[index, column] = None
    return frame


def _apply_filters(df: pd.DataFrame, filters: Optional[Dict[str, List[str]]]) -> pd.DataFrame:
    """Fix 72: Apply global slicer filters to DataFrame."""
    if not filters:
        return df
    for col, values in filters.items():
        if col in df.columns and values:
            df = df[df[col].isin(values)]
    return df


def _natural_sort_key(value: Any):
    """「1月 2月 … 10月」按数字顺序排, 不按字符串排成「10月 11月 12月 1月」。"""
    import re
    text = str(value)
    return [(0, int(part), "") if part.isdigit() else (1, 0, part) for part in re.split(r"(\d+)", text) if part]


def _extract_dimensions(df: pd.DataFrame) -> List[Dict[str, Any]]:
    """Fix 72: Extract available filter dimensions from DataFrame."""
    dims = []
    mapping = dashboard_service.normalizer.detect_columns(df.columns.tolist(), df)
    measures = set(mapping.actual_cols + mapping.budget_cols + mapping.last_year_cols
                   + mapping.last_year_budget_cols)
    # Amount-aware normalization may prefer revenue/cost over the ratio columns.
    # Those original percentage measures are still measures, never dimensions.
    measures.update(col for col in df.columns
                    if any(key in str(col).lower() for key in
                           dashboard_service.normalizer.RATIO_HEADER_KEYWORDS + ('ratio', 'margin', 'percent')))
    for col in df.columns:
        if col in measures:
            continue
        if df[col].dtype == 'object':
            unique_vals = df[col].dropna().unique().tolist()
            if 2 <= len(unique_vals) <= 20:
                dims.append({"name": col, "values": sorted(unique_vals, key=_natural_sort_key)})
    return dims


@router.post("/generate")
async def generate_chart(request: FinancialDashboardRequest, http_request: Request):
    """Generate a single financial chart or all charts."""
    try:
        factory_id = _resolve_factory_id(http_request, request.factory_id)
        df = await _load_data(request, factory_id, http_request)
        df = _apply_filters(df, request.filters)  # Fix 72
        result = dashboard_service.generate_chart(
            chart_type=request.chart_type,
            raw_data=df,
            year=request.year,
            period_type=request.period_type,
            start_month=request.start_month,
            end_month=request.end_month,
            sheet_name=request.sheet_name,
            filters=request.filters,
        )
        return _sanitize_for_json(result)
    except DataSourceError as e:
        return e.response()
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Chart generation failed: {e}", exc_info=True)
        return {"success": False, "error": "图表生成失败，请检查数据格式后重试"}


@router.post("/batch")
async def batch_generate(request: FinancialDashboardRequest, http_request: Request):
    """Generate all available charts (batch mode)."""
    try:
        request.chart_type = "all"
        factory_id = _resolve_factory_id(http_request, request.factory_id)
        df = await _load_data(request, factory_id, http_request)
        # Fix 72: Extract dimensions before filtering
        dimensions = _extract_dimensions(df)
        # Fix 72: Apply slicer filters
        df = _apply_filters(df, request.filters)
        result = dashboard_service.generate_dashboard(
            raw_data=df,
            year=request.year,
            period_type=request.period_type,
            start_month=request.start_month,
            end_month=request.end_month,
            sheet_name=request.sheet_name,
            filters=request.filters,
        )
        # Fix 72: Attach available dimensions to response
        if isinstance(result, dict):
            result["availableDimensions"] = dimensions
        return _sanitize_for_json(result)
    except DataSourceError as e:
        return e.response()
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Batch generation failed: {e}", exc_info=True)
        return {"success": False, "error": "批量生成失败，请检查数据格式后重试"}


@router.post("/analyze")
async def analyze_chart(request: AnalyzeRequest):
    """AI analysis for a single chart, with optional follow-up conversation."""
    try:
        from common.llm_router import call_chain, SLOT
        from common.llm_metrics import llm_caller_context

        system_msg = "你是服务于食品加工企业CFO的资深财务分析师。请用简洁的中文回复。"

        if request.follow_up_question and request.previous_analysis:
            # Follow-up: multi-turn conversation with history
            messages = [
                {"role": "system", "content": system_msg},
                {"role": "user", "content": (
                    f"请对以下财务数据进行分析：\n"
                    f"图表类型：{request.chart_type}\n"
                    f"数据摘要：{request.analysis_context}"
                )},
                {"role": "assistant", "content": request.previous_analysis},
                {"role": "user", "content": request.follow_up_question},
            ]
        else:
            # Initial analysis
            prompt = (
                f"请对以下财务数据进行专业分析：\n\n"
                f"图表类型：{request.chart_type}\n"
                f"数据摘要：{request.analysis_context}\n\n"
                f"请提供：\n"
                f"1. 核心发现（3-5条）\n"
                f"2. 风险预警（如有）\n"
                f"3. 改进建议（2-3条）\n\n"
                f"要求：结合食品加工行业特点，数据驱动，具体可操作。"
            )
            messages = [
                {"role": "system", "content": system_msg + "请用要点列表格式。"},
                {"role": "user", "content": prompt},
            ]

        # 走 call_chain(SLOT.INSIGHTS): 免费 fallback 链 + 熔断。原 OpenAI SDK 同步直连
        # (在 async endpoint 内阻塞) 改为 async router 调用。model 由 router 按 slot 注入。
        payload = {
            "messages": messages,
            "max_tokens": 2000,
            "temperature": 0.3,
        }
        with llm_caller_context("financial_dashboard_analyze"):
            data = await call_chain(SLOT.INSIGHTS, payload, timeout=60.0)

        analysis_text = data["choices"][0]["message"]["content"]
        return {
            "success": True,
            "chartType": request.chart_type,
            "analysis": analysis_text,
        }
    except Exception as e:
        logger.error(f"Analysis failed: {e}", exc_info=True)
        return {"success": False, "error": "AI 分析暂时不可用，请稍后重试"}


@router.get("/templates")
async def list_templates():
    """List all available chart templates."""
    return {
        "success": True,
        "templates": dashboard_service.list_templates(),
    }


@router.post("/export-ppt")
async def export_ppt(request: PPTExportRequest):
    """Export financial dashboard to PPTX file."""
    try:
        from smartbi.services.ppt_generator import FinancialPPTGenerator
        generator = FinancialPPTGenerator()
        pptx_bytes = generator.generate(
            chart_images=request.chart_images,
            headline=request.headline,
            findings=request.findings,
            analysis_results=request.analysis_results,
            company_name=request.company_name,
            year=request.year,
            period_type=request.period_type,
            start_month=request.start_month,
            end_month=request.end_month,
            kpi_summary=request.kpi_summary,
            template=request.template,
            chart_titles=request.chart_titles,
            period_label=request.period_label,
            report_title=request.report_title,
        )
        from urllib.parse import quote
        filename = f"财务分析报告_{request.year}年{request.start_month}-{request.end_month}月.pptx"
        filename_ascii = f"report_{request.year}_{request.start_month}-{request.end_month}.pptx"
        filename_encoded = quote(filename)
        return StreamingResponse(
            io.BytesIO(pptx_bytes),
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename_ascii}\"; filename*=UTF-8''{filename_encoded}",  # noqa: E501
            },
        )
    except Exception as e:
        logger.error(f"PPT export failed: {e}", exc_info=True)
        return {"success": False, "error": "PPT 导出失败，请稍后重试"}


class ExcelExportRequest(BaseModel):
    _export_limits = root_validator(pre=True, allow_reuse=True)(validate_export)
    chart_titles: Dict[str, str] = Field(default_factory=dict)
    period_label: Optional[str] = None
    report_title: Optional[str] = None
    charts: List[Dict[str, Any]] = Field(default_factory=list)  # Chart results from batch
    analysis_results: Dict[str, str] = Field(default_factory=dict)
    # 看板上的页首结论与规则结论 (/batch 原样回传): 导出报告里要和页面上读到的是同一组话
    headline: Optional[Dict[str, Any]] = None
    findings: Optional[Dict[str, List[str]]] = None
    company_name: str = "白垩纪科技"
    year: int = 2026
    period_type: str = "year"
    start_month: int = 1
    end_month: int = 12


@router.post("/export-excel")
async def export_excel(request: ExcelExportRequest):
    """Export financial dashboard data to Excel with multiple sheets."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        if request.report_title:
            info = wb.create_sheet('报告信息')
            info.append(['报告', request.report_title])
            info.append(['工厂', request.company_name])
            info.append(['期间', request.period_label or f'{request.year}年{request.start_month}–{request.end_month}月'])
            info.append(['说明', '仅包含页面当前显示且选中的图表；逐图来源、单位与规则分析见明细及分析表'])

        # Header styles
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="1B65A8", end_color="1B65A8", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(  # noqa: F841
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # Sheet 1: 分析结论 (有页首结论或规则结论时)
        conclusion_rows = _conclusion_rows(request.headline, request.findings)
        if conclusion_rows:
            ws_conc = wb.active
            ws_conc.title = "分析结论"
            ws_conc.append(["类别", "内容"])
            for cell in ws_conc[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
            for row in conclusion_rows:
                ws_conc.append(row)
            ws_conc.column_dimensions["A"].width = 14
            ws_conc.column_dimensions["B"].width = 90
            ws_summary = wb.create_sheet(title="KPI汇总")
        else:
            ws_summary = wb.active
            ws_summary.title = "KPI汇总"
        ws_summary.append(["图表类型", "指标", "数值", "单位", "趋势"])
        for cell in ws_summary[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        for chart in request.charts:
            chart_type = chart.get("chartType", "")
            title = chart.get("title", chart_type)
            for kpi in chart.get("kpis", []):
                ws_summary.append([
                    title,
                    kpi.get("label", ""),
                    kpi.get("value", ""),
                    kpi.get("unit", ""),
                    {"up": "↑", "down": "↓", "flat": "→"}.get(kpi.get("trend", ""), ""),
                ])

        # Auto-width for summary
        for col in ws_summary.columns:
            max_len = max(len(str(cell.value or "")) for cell in col) + 2
            ws_summary.column_dimensions[col[0].column_letter].width = min(max_len, 30)

        # Sheet per chart with tableData
        from smartbi.services.ppt_generator import CHART_DISPLAY_NAMES
        for chart in request.charts:
            chart_type = chart.get("chartType", "")
            table_data = chart.get("tableData")
            if not table_data:
                continue

            display_name = chart.get('title') or request.chart_titles.get(chart_type) or CHART_DISPLAY_NAMES.get(chart_type, chart_type)
            # Sheet name max 31 chars
            import re
            sheet_name = re.sub(r'[\\/*?:\[\]]', ' ', str(display_name)).strip()[:31] or '图表明细'
            ws = wb.create_sheet(title=sheet_name)
            if request.report_title:
                for label, value in [('来源',chart.get('source','')),('期间',chart.get('period',request.period_label or '')),
                                     ('单位',chart.get('unit','')),('规则依据',chart.get('reason',''))]:
                    ws.append([label,value])

            # 明细表里的金额按图表显示倍率缩放过 (万/亿), 必须写明单位, 否则拿去求和对账会错一万倍
            unit_note = _table_unit_note(chart)
            if unit_note:
                ws.append([unit_note])
            header_row = ws.max_row + 1 if unit_note or request.report_title else 1
            ws.append(table_data.get("headers", []))
            for cell in ws[header_row]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align

            # 构建器的明细行有两种形状: {"label", "values"} 或整行列表
            for row in table_data.get("rows", []):
                if isinstance(row, dict):
                    ws.append([row.get("label", "")] + list(row.get("values") or []))
                elif isinstance(row, (list, tuple)):
                    ws.append(list(row))
                else:
                    raise ValueError(f"无法识别的明细行格式: {type(row).__name__}")

            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col) + 2
                ws.column_dimensions[col[0].column_letter].width = min(max_len, 25)

        # Sheet: AI Analysis
        if request.analysis_results:
            ws_ai = wb.create_sheet(title="逐图规则与来源" if request.report_title else "AI分析")
            ws_ai.append(["图表", "分析内容"])
            for cell in ws_ai[1]:
                cell.font = header_font
                cell.fill = header_fill
            for chart_type, analysis in request.analysis_results.items():
                display_name = request.chart_titles.get(chart_type) or CHART_DISPLAY_NAMES.get(chart_type, chart_type)
                ws_ai.append([display_name, analysis if request.report_title else analysis[:5000]])
            ws_ai.column_dimensions["A"].width = 20
            ws_ai.column_dimensions["B"].width = 80

        if request.report_title:
            # Treat uploaded strings as text in every sheet, including metadata.
            import re
            for sheet in wb:
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value, str) and re.match(r'^\s*[=+@-]', cell.value):
                            cell.value = "'" + cell.value
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        from urllib.parse import quote
        filename = f"财务分析数据_{request.year}年{request.start_month}-{request.end_month}月.xlsx"
        filename_ascii = f"financial_data_{request.year}_{request.start_month}-{request.end_month}.xlsx"

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename_ascii}\"; filename*=UTF-8''{quote(filename)}",
            },
        )
    except ImportError:
        return _export_error("服务器未安装 openpyxl，请联系管理员")
    except Exception as e:
        logger.error(f"Excel export failed: {e}", exc_info=True)
        return _export_error("Excel 导出失败，请稍后重试")


def _export_error(message: str) -> JSONResponse:
    # 非 200: 前端据此不会把错误 JSON 当成 .xlsx 下载并提示成功
    return JSONResponse(status_code=500, content={"success": False, "message": message, "error": message})


def _table_unit_note(chart: Dict[str, Any]) -> Optional[str]:
    """图表明细的金额单位说明 (元 / 万元 / 亿元)。"""
    scale = (chart.get("metadata") or {}).get("scale")
    suffix = scale.get("suffix") if isinstance(scale, dict) else scale
    if suffix in ("万", "亿"):
        return f"金额单位：{suffix}元（百分比列除外）"
    if suffix == "":
        return "金额单位：元（百分比列除外）"
    return None


class PDFExportRequest(BaseModel):
    _export_limits = root_validator(pre=True, allow_reuse=True)(validate_export)
    chart_titles: Dict[str, str] = Field(default_factory=dict)
    period_label: Optional[str] = None
    report_title: Optional[str] = None
    chart_images: Dict[str, str] = Field(default_factory=dict)  # {chartType: base64PNG}
    headline: Optional[Dict[str, Any]] = None
    findings: Optional[Dict[str, List[str]]] = None
    analysis_results: Dict[str, str] = Field(default_factory=dict)
    company_name: str = "白垩纪科技"
    year: int = 2026
    period_type: str = "year"
    start_month: int = 1
    end_month: int = 12
    kpi_summary: Optional[Dict] = None


@router.post("/export-pdf")
async def export_pdf(request: PDFExportRequest):
    """Export financial dashboard to PDF with chart images and analysis."""
    try:
        from smartbi.services.pdf_generator import FinancialPDFGenerator
        generator = FinancialPDFGenerator()
        pdf_bytes = generator.generate(
            chart_images=request.chart_images,
            headline=request.headline,
            findings=request.findings,
            analysis_results=request.analysis_results,
            company_name=request.company_name,
            year=request.year,
            period_type=request.period_type,
            start_month=request.start_month,
            end_month=request.end_month,
            kpi_summary=request.kpi_summary,
            chart_titles=request.chart_titles,
            period_label=request.period_label,
            report_title=request.report_title,
        )

        from urllib.parse import quote
        filename = f"财务分析报告_{request.year}年{request.start_month}-{request.end_month}月.pdf"
        filename_ascii = f"report_{request.year}_{request.start_month}-{request.end_month}.pdf"

        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={
                "Content-Disposition": f"attachment; filename=\"{filename_ascii}\"; filename*=UTF-8''{quote(filename)}",
            },
        )
    except ImportError as e:
        logger.error(f"PDF export import failed: {e}")
        return {"success": False, "error": "PDF 导出依赖未安装，请联系管理员"}
    except Exception as e:
        logger.error(f"PDF export failed: {e}", exc_info=True)
        return {"success": False, "error": "PDF 导出失败，请稍后重试"}


# Same authenticated namespace; query-only pages never modify source records.
from smartbi.api.factory_chart_details import router as _factory_details_router
router.include_router(_factory_details_router)
