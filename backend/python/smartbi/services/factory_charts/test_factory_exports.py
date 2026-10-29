"""Read actual exported documents, not only successful HTTP responses."""
import base64
import io

from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
import openpyxl
from pptx import Presentation
from pypdf import PdfReader

from smartbi.api import financial_dashboard as api
from smartbi.services.ppt_generator import _wrap_for_ppt


def client():
    app=FastAPI()
    @app.middleware('http')
    async def identity(request,call_next):
        request.state.factory_id='F006'; request.state.role='factory_super_admin'
        return await call_next(request)
    app.include_router(api.router)
    return TestClient(app)


def payload():
    output=io.BytesIO(); Image.new('RGB',(400,150),'white').save(output,format='PNG')
    image='data:image/png;base64,'+base64.b64encode(output.getvalue()).decode()
    return {'year':2027,'period_type':'month_range','start_month':2,'end_month':3,
        'period_label':'2027年2–3月','report_title':'工厂经营分析报告','company_name':'F006',
        'chart_titles':{'one':'产品单位成本/甲','two':'检验合格率'},
        'chart_images':{'one':image,'two':image},
        'charts':[{'chartType':'bar_horizontal__one','title':'产品单位成本/甲',
            'source':'F006上传报表','period':'2027年2–3月','unit':'元/kg','reason':'总成本÷总产量',
            'kpis':[{'label':'累计单位成本','value':'2.64','unit':'元/kg'}],
            'tableData':{'headers':['SKU','单位成本（元/kg）'], 'rows':[['产品A',2.64]]}},
            {'chartType':'bar__two','title':'检验合格率','unit':'%',
            'tableData':{'headers':['批次','合格率（%）'], 'rows':[{'label':'批次B','values':[91.79]}]}}],
        'analysis_results':{'one':'累计单位成本：2.64元/kg\n来源：F006上传报表\n期间：2027年2–3月\n'+'逐图规则完整保留\n'*200+'末尾事实校验',
                            'two':'检验合格率：91.79%\n合格数÷检验数；不是生产出成率'}}


def test_factory_excel_keeps_titles_period_both_row_shapes_units_and_facts():
    response=client().post('/export-excel',json=payload())
    assert response.status_code==200 and 'spreadsheetml' in response.headers['content-type']
    book=openpyxl.load_workbook(io.BytesIO(response.content),data_only=False)
    assert '报告信息' in book.sheetnames and '产品单位成本 甲' in book.sheetnames and '检验合格率' in book.sheetnames
    cells=[[cell.value for cell in row] for sheet in book for row in sheet]
    assert ['期间','2027年2–3月'] in cells
    assert ['产品A',2.64] in cells and ['批次B',91.79] in cells
    assert any('2.64' in str(row) and '元/kg' in str(row) for row in cells)
    assert any('末尾事实校验' in str(row) for row in cells)


def test_factory_excel_uploaded_formulas_are_literal_cells():
    data=payload(); data['charts'][0]['tableData']['rows'][0][0]=' \t=HYPERLINK("https://invalid")'
    data['charts'][0]['source']='=NOT_A_FORMULA'
    response=client().post('/export-excel',json=data)
    book=openpyxl.load_workbook(io.BytesIO(response.content),data_only=False)
    sheet=book['产品单位成本 甲']
    assert any(cell.value=='\'=NOT_A_FORMULA' and cell.data_type=='s' for row in sheet for cell in row)
    assert any(str(cell.value).startswith("' \t=HYPERLINK") and cell.data_type=='s' for row in sheet for cell in row)


def test_factory_excel_formula_guard_also_applies_without_report_title():
    data=payload()
    data['report_title']=None
    data['charts'][0]['tableData']['rows'][0][0]='=HYPERLINK("https://invalid")'
    response=client().post('/export-excel',json=data)
    book=openpyxl.load_workbook(io.BytesIO(response.content),data_only=False)
    sheet=book['产品单位成本 甲']
    assert any(cell.value=="'=HYPERLINK(\"https://invalid\")" and cell.data_type=='s' for row in sheet for cell in row)


def test_factory_excel_print_layout_wraps_provenance_without_clipping_columns():
    response=client().post('/export-excel',json=payload())
    book=openpyxl.load_workbook(io.BytesIO(response.content),data_only=False)
    sheet=book['产品单位成本 甲']
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.orientation == 'landscape'
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert sheet.print_area == "'产品单位成本 甲'!$A$1:$B$6"
    assert sheet['B1'].alignment.wrap_text is True
    assert sheet['B1'].value == 'F006上传报表'


def test_factory_excel_long_analysis_is_split_into_printable_rows():
    response=client().post('/export-excel',json=payload())
    book=openpyxl.load_workbook(io.BytesIO(response.content),data_only=False)
    sheet=book['逐图规则与来源']
    values='\n'.join(str(cell.value or '') for row in sheet.iter_rows() for cell in row)
    assert sheet.max_row > 5
    assert '末尾事实校验' in values
    assert '逐图规则完整保留' in values
    assert max((row.height or 0) for row in sheet.row_dimensions.values()) <= 120


def test_factory_ppt_actual_titles_period_images_kpi_and_last_fact_survive():
    response=client().post('/export-ppt',json=payload())
    assert response.status_code==200 and response.content.startswith(b'PK')
    ppt=Presentation(io.BytesIO(response.content))
    text='\n'.join(shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame)
    for value in ['工厂经营分析报告','产品单位成本/甲','检验合格率','2027年2–3月','2.64元/kg','91.79%','末尾事实校验']:
        assert value in text
    assert sum(shape.shape_type==13 for slide in ppt.slides for shape in slide.shapes)==2
    assert 'AI 分析' not in text


def test_factory_ppt_keeps_ascii_status_tokens_intact_when_paginating():
    data=payload()
    data['analysis_results']['one'] = ('需要关注：发货必须通过sales_order_item_id精确关联；当前库存仅为AVAILABLE快照，不倒推历史\n' * 40)
    response=client().post('/export-ppt',json=data)
    ppt=Presentation(io.BytesIO(response.content))
    texts=[shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame]
    assert any('AVAILABLE快照' in value for value in texts)
    assert not any(value.strip() in {'AVAIL', 'ABLE快照', 'LE快照'} for value in texts)


def test_factory_ppt_balances_single_cjk_tail_without_losing_fact_text():
    data=payload()
    fact='需要关注：仅统计订单与发货关联校验通过的订单行；缺少关系或单位不一致的行不计入，不代表全部销售'
    data['analysis_results']['one'] = fact
    response=client().post('/export-ppt',json=data)
    assert response.status_code==200 and response.content.startswith(b'PK')
    ppt=Presentation(io.BytesIO(response.content))
    assert len(ppt.slides) == 6
    texts=[shape.text for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame]
    fact_shapes=[value for value in texts if fact[:12] in value]
    assert fact_shapes
    assert any(''.join(value.splitlines()) == fact for value in fact_shapes)
    assert not any(line.strip() == '售' for value in fact_shapes for line in value.splitlines())
    assert _wrap_for_ppt('合格\n售') == ['合格', '售']


def test_factory_pdf_actual_titles_period_images_kpi_and_last_fact_survive():
    response=client().post('/export-pdf',json=payload())
    assert response.status_code==200 and response.content.startswith(b'%PDF')
    pdf=PdfReader(io.BytesIO(response.content)); text='\n'.join(page.extract_text() or '' for page in pdf.pages)
    for value in ['工厂经营分析报告','产品单位成本/甲','检验合格率','2027年2–3月','2.64元/kg','91.79%','末尾事实校验']:
        assert value in text
    assert sum(len(page.images) for page in pdf.pages)==2
    assert 'AI 分析' not in text
