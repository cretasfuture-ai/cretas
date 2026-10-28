"""Exercise the real export endpoints and read the files the user downloads."""
import io

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader

from smartbi.api.financial_dashboard import router


HEADLINE = '收入 1,234.00，毛利率 28.0%'
FINDINGS = {
    'good': ['收入比上年增加 234.00'],
    'attention': ['成本 < 预算 & 费用 > 上年；保留 <b>原表</b> 字样'],
    'suggestions': ['复盘2月：毛利率下降 2.0 个百分点'],
}


def export(kind, **payload):
    app = FastAPI()
    @app.middleware('http')
    async def authenticated_factory(request,call_next):
        request.state.factory_id='F006'
        request.state.role='factory_super_admin'
        return await call_next(request)
    app.include_router(router)
    response = TestClient(app).post(f'/export-{kind}', json=payload)
    assert response.status_code == 200
    if kind == 'ppt':
        assert 'presentationml' in response.headers['content-type']
        return ['\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                for slide in Presentation(io.BytesIO(response.content)).slides]
    assert response.headers['content-type'] == 'application/pdf'
    return [page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages]


@pytest.mark.parametrize('kind', ['ppt', 'pdf'])
def test_conclusions_precede_charts_and_preserve_every_visible_fact(kind):
    pages = export(kind, headline={'text': HEADLINE}, findings=FINDINGS,
                   chart_images={'budget_achievement': ''},
                   analysis_results={'budget_achievement': 'AI 补充内容'})
    assert '财务分析报告' in pages[0]
    conclusion = pages[1]
    for expected in ['分析结论', '页首结论', HEADLINE, '做得好的', '需要关注', '建议']:
        assert expected in conclusion
    for lines in FINDINGS.values():
        assert all(line in conclusion for line in lines)
    assert '预算完成情况' in pages[2]
    assert 'AI 补充内容' in pages[2]
    assert 'AI 补充内容' not in conclusion


@pytest.mark.parametrize('kind', ['ppt', 'pdf'])
@pytest.mark.parametrize('payload', [{}, {'headline': None, 'findings': None},
                                   {'headline': {'text': ' '}, 'findings': {'good': [' ']}}])
def test_no_rule_content_does_not_create_an_empty_conclusion_page(kind, payload):
    pages = export(kind, chart_images={'budget_achievement': ''}, **payload)
    assert len(pages) == 3
    assert '预算完成情况' in pages[1]
    assert all('分析结论' not in page for page in pages)


@pytest.mark.parametrize('kind', ['ppt', 'pdf'])
def test_headline_only_and_findings_only_remain_useful(kind):
    assert HEADLINE in export(kind, headline={'text': HEADLINE})[1]
    assert FINDINGS['good'][0] in export(kind, findings={'good': FINDINGS['good']})[1]


@pytest.mark.parametrize('kind', ['ppt', 'pdf'])
def test_long_conclusions_paginate_without_losing_the_last_fact(kind):
    findings = {'attention': [f'第{i}项成本增加 {i}.00，需要关注原料成本与费用的月度变化'
                              for i in range(35)]}
    pages = export(kind, findings=findings, chart_images={'budget_achievement': ''})
    assert len(pages) > 4
    text = ''.join(pages[1:-2]).replace('\n', '')
    for line in findings['attention']:
        assert line in text
